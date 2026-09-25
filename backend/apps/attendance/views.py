import csv
from io import BytesIO

import qrcode
import qrcode.image.svg
from django.conf import settings
from django.contrib.auth import get_user_model
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework import mixins, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response
from rest_framework.throttling import ScopedRateThrottle
from rest_framework.views import APIView

from apps.academics.models import CourseOffering, Enrollment
from apps.accounts.permissions import IsStudent, Requires
from apps.core import spreadsheets
from apps.core.services import audit, client_ip

from . import services, uploads
from .models import AttendanceCorrection, AttendanceRecord, AttendanceSession
from .serializers import (
    CheckInSerializer,
    CorrectionRequestSerializer,
    CorrectionSerializer,
    DecisionSerializer,
    MarkSerializer,
    RecordSerializer,
    SessionSerializer,
)

User = get_user_model()
ATTENDED = Q(records__status__in=AttendanceRecord.ATTENDED)


def can_view_offering(user, offering):
    """The course lecturer, or staff with attendance.view over the course's department."""
    if offering.lecturer_id == user.id:
        return True
    department = offering.course.department
    return user.has_permission("attendance.view") and user.can_act_on_department(
        "attendance.view", department.pk, department.faculty_id
    )


# --- Lecturers ----------------------------------------------------------------------------------


class SessionViewSet(
    mixins.ListModelMixin,
    mixins.RetrieveModelMixin,
    mixins.CreateModelMixin,
    mixins.DestroyModelMixin,
    viewsets.GenericViewSet,
):
    """Attendance sessions for an offering (?offering=<id>). Lecturers create and run them."""

    serializer_class = SessionSerializer
    pagination_class = None

    def get_queryset(self):
        qs = AttendanceSession.objects.select_related("offering__course__department", "offering__semester").annotate(
            attended_count=Count("records", filter=ATTENDED), record_count=Count("records")
        )
        user = self.request.user
        if not user.is_super_admin:
            qs = qs.filter(
                Q(offering__lecturer=user)
                | (
                    user.scope_filter("attendance.view", "offering__course__department")
                    if user.has_permission("attendance.view")
                    else Q(pk__in=[])
                )
            )
        offering = self.request.query_params.get("offering")
        qs = qs.order_by("-date", "-start_time")  # Meta.ordering doesn't apply to the counting query
        return qs.filter(offering_id=offering) if offering else qs

    def perform_create(self, serializer):
        offering = serializer.validated_data["offering"]
        services.ensure_lecturer(self.request.user, offering)
        session = serializer.save(created_by=self.request.user)
        audit(self.request, "attendance.session_create", session, f"Created attendance session {session}")

    def perform_destroy(self, session):
        services.ensure_lecturer(self.request.user, session.offering)
        if session.status != AttendanceSession.Status.SCHEDULED:
            raise ValidationError("Only sessions that haven't started can be deleted.")
        session.delete()

    def retrieve(self, request, *args, **kwargs):
        """The session with every registered student and their record (if any)."""
        session = self.get_object()
        records = {r.student_id: r for r in session.records.select_related("student").prefetch_related("corrections")}
        roster = []
        for student in services.registered_students(session.offering):
            record = records.pop(student.pk, None)
            roster.append(
                {
                    "student": student.pk,
                    "student_name": student.get_full_name(),
                    "matric_number": student.university_id,
                    "record": RecordSerializer(record).data if record else None,
                }
            )
        # Students who dropped the course after being marked still appear, for the record.
        roster += [
            {
                "student": r.student_id,
                "student_name": r.student.get_full_name(),
                "matric_number": r.student.university_id,
                "record": RecordSerializer(r).data,
            }
            for r in records.values()
        ]
        can_mark = session.offering.lecturer_id == request.user.id or request.user.is_super_admin
        return Response({**SessionSerializer(session).data, "roster": roster, "can_mark": can_mark})

    @action(detail=True, methods=["post"])
    def open(self, request, pk=None):
        session = services.open_session(self.get_object(), request.user)
        audit(request, "attendance.session_open", session, f"Opened check-in for {session}")
        return Response(SessionSerializer(session).data)

    @action(detail=True, methods=["post"])
    def close(self, request, pk=None):
        session = services.close_session(self.get_object(), request.user)
        audit(request, "attendance.session_close", session, f"Closed {session}")
        return Response(SessionSerializer(session).data)

    @action(detail=True, methods=["get"])
    def qr(self, request, pk=None):
        """The current QR code and 6-digit code, for the lecturer's screen (polled)."""
        session = self.get_object()
        services.ensure_lecturer(request.user, session.offering)
        if not session.is_accepting_checkins:
            return Response({"accepting": False, "status": session.status})
        token, code, seconds_left = services.current_codes(session)
        url = request.build_absolute_uri(f"/portal/attend?s={session.pk}&t={token}")
        svg = qrcode.make(url, image_factory=qrcode.image.svg.SvgPathImage, box_size=12, border=2)
        return Response(
            {
                "accepting": True,
                "code": code,
                "url": url,
                "seconds_left": seconds_left,
                "svg": svg.to_string(encoding="unicode"),
                "checkin_closes_at": session.checkin_closes_at,
                "checked_in": session.records.filter(status__in=AttendanceRecord.ATTENDED).count(),
                "registered": services.registered_students(session.offering).count(),
            }
        )

    @action(detail=True, methods=["post"])
    def mark(self, request, pk=None):
        session = self.get_object()
        serializer = MarkSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        student = get_object_or_404(User, pk=serializer.validated_data["student"])
        record = services.mark(session, student, serializer.validated_data["status"], request.user)
        return Response(RecordSerializer(record).data)


class RecordCorrectionView(APIView):
    """Lecturer asks to change a record in a closed session (needs HOD approval)."""

    def post(self, request, pk):
        record = get_object_or_404(AttendanceRecord.objects.select_related("session__offering"), pk=pk)
        serializer = CorrectionRequestSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        correction = services.request_correction(
            record, serializer.validated_data["to_status"], serializer.validated_data["reason"], request.user
        )
        audit(request, "attendance.correction_request", correction, f"Requested correction for {record}")
        return Response(CorrectionSerializer(correction).data, status=status.HTTP_201_CREATED)


class OfferingStatsView(APIView):
    """Per-student attendance for a course offering, with a summary."""

    def get(self, request, pk):
        offering = get_object_or_404(CourseOffering.objects.select_related("course__department"), pk=pk)
        if not can_view_offering(request.user, offering):
            raise PermissionDenied("You can't view attendance for this course.")
        students = list(services.registered_students(offering))
        stats = services.offering_student_stats(offering, students)
        minimum = settings.ATTENDANCE_MIN_PERCENT
        rows = [
            {
                "student": s.pk,
                "student_name": s.get_full_name(),
                "matric_number": s.university_id,
                **stats[s.pk],
                "at_risk": stats[s.pk]["percent"] is not None and stats[s.pk]["percent"] < minimum,
            }
            for s in students
        ]
        percents = [r["percent"] for r in rows if r["percent"] is not None]
        held = rows[0]["held"] if rows else offering.attendance_sessions.filter(status="closed").count()
        return Response(
            {
                "course_code": offering.course.code,
                "course_title": offering.course.title,
                "minimum_percent": minimum,
                "summary": {
                    "sessions_held": held,
                    "average_percent": round(sum(percents) / len(percents), 1) if percents else None,
                    "at_risk": sum(1 for r in rows if r["at_risk"]),
                    "students": len(rows),
                },
                "students": rows,
            }
        )


# --- Students -----------------------------------------------------------------------------------


class CheckInView(APIView):
    """Self check-in by QR link (session + token) or by typing the 6-digit code."""

    permission_classes = [IsStudent]
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = "checkin"

    def post(self, request):
        serializer = CheckInSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        if data.get("session"):
            session = get_object_or_404(
                AttendanceSession.objects.select_related("offering__course"), pk=data["session"]
            )
        else:
            session = self._session_for_code(request.user, data["code"])
        record, created = services.check_in(
            request.user,
            session,
            token=data["token"],
            code=data["code"],
            device_id=data["device_id"],
            ip_address=client_ip(request),
        )
        message = f"You're checked in for {session.offering.course.code}" + (
            " (late)." if record.status == AttendanceRecord.Status.LATE else "."
        )
        if not created:
            message = f"You're already marked {record.get_status_display().lower()} for {session.offering.course.code}."
        return Response(
            {
                "created": created,
                "message": message,
                "course_code": session.offering.course.code,
                "status": record.status,
                "status_label": record.get_status_display(),
            },
            status=status.HTTP_201_CREATED if created else status.HTTP_200_OK,
        )

    @staticmethod
    def _session_for_code(student, code):
        """Find which of the student's open classes a typed code belongs to."""
        sessions = AttendanceSession.objects.filter(
            status=AttendanceSession.Status.OPEN,
            offering__enrollments__student=student,
            offering__enrollments__status=Enrollment.Status.REGISTERED,
        ).select_related("offering__course")
        for session in sessions:
            if services.code_is_valid(session, code=code):
                return session
        raise ValidationError("That code doesn't match an open class of yours. Check it and try again.")


class MyAttendanceView(APIView):
    """A student's attendance per course this semester, with full history."""

    permission_classes = [IsStudent]

    def get(self, request):
        return Response(services.student_attendance(request.user))


# --- Administrators: reports and corrections ----------------------------------------------------


REPORT_FILTERS = {
    "semester": "session__offering__semester",
    "faculty": "session__offering__course__department__faculty",
    "department": "session__offering__course__department",
    "offering": "session__offering",
    "student": "student",
}


def filtered_records(request):
    user = request.user
    records = AttendanceRecord.objects.filter(
        user.scope_filter("attendance.view", "session__offering__course__department")
    )
    for param, lookup in REPORT_FILTERS.items():
        value = request.query_params.get(param)
        if value:
            records = records.filter(**{lookup: value})
    search = request.query_params.get("search", "").strip()
    if search:
        records = records.filter(
            Q(student__university_id__icontains=search)
            | Q(student__first_name__icontains=search)
            | Q(student__last_name__icontains=search)
        )
    return records


def report_data(request):
    rows = services.report_rows(filtered_records(request))
    if request.query_params.get("at_risk") == "true":
        rows = [r for r in rows if r["at_risk"]]
    return rows


class ReportView(APIView):
    permission_classes = [Requires("attendance.view")]

    def get(self, request):
        rows = report_data(request)
        percents = [r["percent"] for r in rows if r["percent"] is not None]
        return Response(
            {
                "minimum_percent": settings.ATTENDANCE_MIN_PERCENT,
                "summary": {
                    "rows": len(rows),
                    "students": len({r["student_id"] for r in rows}),
                    "courses": len({r["offering_id"] for r in rows}),
                    "average_percent": round(sum(percents) / len(percents), 1) if percents else None,
                    "at_risk": sum(1 for r in rows if r["at_risk"]),
                    "absences": sum(r["absent"] for r in rows),
                },
                "rows": rows[:500],
                "truncated": len(rows) > 500,
            }
        )


EXPORT_COLUMNS = [
    ("Matric no.", "matric_number"),
    ("Student", "student_name"),
    ("Course", "course_code"),
    ("Course title", "course_title"),
    ("Department", "department"),
    ("Sessions held", "held"),
    ("Present", "present"),
    ("Late", "late"),
    ("Excused", "excused"),
    ("Absent", "absent"),
    ("Attendance %", "percent"),
    ("Below minimum", "at_risk"),
]


class ReportExportView(APIView):
    """The report as CSV (?file=csv) or Excel (?file=xlsx)."""

    permission_classes = [Requires("attendance.view")]

    def get(self, request):
        rows = report_data(request)
        stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
        audit(request, "attendance.export", None, f"Exported attendance report ({len(rows)} rows)")
        if request.query_params.get("file") == "xlsx":
            return self._xlsx(rows, f"attendance-{stamp}.xlsx")
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="attendance-{stamp}.csv"'
        response.write("﻿")  # BOM so Excel opens UTF-8 names correctly
        writer = csv.writer(response)
        writer.writerow([label for label, _ in EXPORT_COLUMNS])
        for r in rows:
            writer.writerow([("Yes" if r[key] else "No") if key == "at_risk" else r[key] for _, key in EXPORT_COLUMNS])
        return response

    @staticmethod
    def _xlsx(rows, filename):
        book = Workbook()
        sheet = book.active
        sheet.title = "Attendance"
        sheet.append([label for label, _ in EXPORT_COLUMNS])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12305A")
        warn = PatternFill("solid", fgColor="FDECEB")
        for r in rows:
            sheet.append([("Yes" if r[key] else "No") if key == "at_risk" else r[key] for _, key in EXPORT_COLUMNS])
            if r["at_risk"]:
                for cell in sheet[sheet.max_row]:
                    cell.fill = warn
        for column, width in zip("ABCDEFGHIJKL", (16, 28, 10, 34, 30, 14, 9, 7, 9, 9, 13, 14), strict=True):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        buffer = BytesIO()
        book.save(buffer)
        response = HttpResponse(
            buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class CorrectionViewSet(mixins.ListModelMixin, viewsets.GenericViewSet):
    """Approvers see corrections for their departments; lecturers see the ones they requested."""

    serializer_class = CorrectionSerializer
    filterset_fields = ["status"]

    def get_queryset(self):
        user = self.request.user
        qs = AttendanceCorrection.objects.select_related(
            "record__student", "record__session__offering__course", "requested_by", "decided_by"
        )
        visible = Q(requested_by=user)
        if user.has_permission("attendance.approve"):
            visible |= user.scope_filter("attendance.approve", "record__session__offering__course__department")
        return qs.filter(visible)

    @action(detail=True, methods=["post"])
    def decide(self, request, pk=None):
        correction = self.get_object()
        serializer = DecisionSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        correction = services.decide_correction(
            correction, serializer.validated_data["approve"], serializer.validated_data["note"].strip(), request.user
        )
        audit(request, "attendance.correction_decide", correction, f"{correction.get_status_display()} {correction}")
        return Response(CorrectionSerializer(correction).data)


# --- Bulk upload of paper registers ---------------------------------------------------------------


def _flag(request, name, default):
    return str(request.data.get(name, default)).lower() in ("1", "true", "yes")


class OfferingUploadView(APIView):
    """POST a CSV/Excel register for a course (multipart `file`; `dry_run=true` to only check it;
    `close=false` to leave the classes open instead of marking everyone else absent)."""

    parser_classes = [MultiPartParser]

    def post(self, request, pk):
        offering = get_object_or_404(CourseOffering.objects.select_related("course", "semester"), pk=pk)
        services.ensure_lecturer(request.user, offering)
        file = request.FILES.get("file")
        if not file:
            raise ValidationError({"file": "Choose a CSV or Excel file."})
        result = uploads.upload(
            offering,
            file,
            request.user,
            request,
            dry_run=_flag(request, "dry_run", False),
            close=_flag(request, "close", True),
        )
        return Response(result, status=status.HTTP_201_CREATED if result["saved"] else status.HTTP_200_OK)


class OfferingUploadTemplateView(APIView):
    """The class list as a register to fill in (?file=csv|xlsx&date=YYYY-MM-DD)."""

    def get(self, request, pk):
        offering = get_object_or_404(CourseOffering.objects.select_related("course"), pk=pk)
        services.ensure_lecturer(request.user, offering)
        day = spreadsheets.parse_date(request.query_params.get("date", ""))
        file_format = "xlsx" if request.query_params.get("file") == "xlsx" else "csv"
        content = uploads.template(offering, day, file_format)
        name = f"{offering.course.code}-register-{(day or timezone.localdate()):%Y-%m-%d}.{file_format}"
        response = HttpResponse(content, content_type=spreadsheets.XLSX if file_format == "xlsx" else "text/csv")
        response["Content-Disposition"] = f'attachment; filename="{name}"'
        return response
