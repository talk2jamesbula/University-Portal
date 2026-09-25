from datetime import timedelta
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.academics.models import Enrollment
from apps.core.models import AuditLog, Notification
from apps.core.testing import (
    client_for,
    make_course,
    make_department,
    make_offering,
    make_programme,
    make_semester,
    make_staff,
    make_student,
)

from ..models import AttendanceRecord, AttendanceSession

Status = AttendanceRecord.Status


def csv_file(*rows, header="date,matric_number,status"):
    return SimpleUploadedFile("register.csv", ("\n".join([header, *rows]) + "\n").encode())


class AttendanceUploadTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        today = timezone.localdate()
        cls.semester = make_semester(start_date=today - timedelta(days=40), end_date=today + timedelta(days=100))
        department = make_department()
        cls.lecturer = make_staff(roles=["lecturer"], department=department)
        cls.offering = make_offering(make_course(department, code="CSC201"), cls.semester, lecturer=cls.lecturer)
        programme = make_programme(department)
        cls.students = [make_student(programme=programme) for _ in range(3)]
        for s in cls.students:
            Enrollment.objects.create(student=s, offering=cls.offering, status=Enrollment.Status.REGISTERED)
        cls.outsider = make_student(programme=programme)
        cls.day = today - timedelta(days=7)

    def setUp(self):
        self.client = client_for(self.lecturer)
        self.url = f"/api/attendance/offerings/{self.offering.pk}/upload/"

    def post(self, file, **extra):
        return self.client.post(self.url, {"file": file, **extra}, format="multipart")

    def matric(self, i):
        return self.students[i].university_id

    def test_check_first_then_save_and_close(self):
        rows = [f"{self.day},{self.matric(0)},P", f"{self.day:%d/%m/%Y},{self.matric(1).lower()},late"]
        preview = self.post(csv_file(*rows), dry_run="true").data
        self.assertFalse(preview["saved"])
        self.assertEqual(preview["sessions"][0]["session_status"], "new")
        self.assertEqual((preview["sessions"][0]["present"], preview["sessions"][0]["late"]), (1, 1))
        self.assertEqual(preview["sessions"][0]["not_in_file"], 1)
        self.assertFalse(AttendanceSession.objects.exists())

        with self.captureOnCommitCallbacks(execute=True):
            response = self.post(csv_file(*rows))
        self.assertEqual(response.status_code, 201, response.data)
        session = AttendanceSession.objects.get()
        self.assertEqual((session.date, session.status), (self.day, AttendanceSession.Status.CLOSED))
        statuses = dict(session.records.values_list("student_id", "status"))
        self.assertEqual(
            statuses,
            {self.students[0].pk: Status.PRESENT, self.students[1].pk: Status.LATE, self.students[2].pk: Status.ABSENT},
        )
        self.assertEqual(session.records.get(student=self.students[0]).method, AttendanceRecord.Method.UPLOAD)
        self.assertTrue(Notification.objects.filter(user=self.students[2], title__contains="absent").exists())
        self.assertTrue(AuditLog.objects.filter(action="attendance.upload").exists())

    def test_several_classes_and_leaving_them_open(self):
        other_day = self.day - timedelta(days=2)
        rows = [
            f"{self.day},{self.matric(0)},present",
            f"{other_day},{self.matric(0)},absent",
            f"{other_day},{self.matric(1)},E",
        ]
        self.assertEqual(self.post(csv_file(*rows), close="false").status_code, 201)
        self.assertEqual(AttendanceSession.objects.count(), 2)
        self.assertFalse(AttendanceSession.objects.filter(status=AttendanceSession.Status.CLOSED).exists())
        self.assertEqual(AttendanceRecord.objects.count(), 3)
        # A second upload updates the open classes rather than duplicating them.
        self.post(csv_file(f"{self.day},{self.matric(0)},late"), close="false")
        self.assertEqual(
            AttendanceRecord.objects.get(session__date=self.day, student=self.students[0]).status, Status.LATE
        )

    def test_every_row_is_checked_and_nothing_saved_on_error(self):
        future = timezone.localdate() + timedelta(days=3)
        rows = [
            f"{self.day},{self.matric(0)},present",
            f"{self.day},{self.matric(0)},absent",  # duplicate
            f"{self.day},{self.outsider.university_id},present",  # not registered
            f"{future},{self.matric(1)},present",  # future
            f"not-a-date,{self.matric(1)},present",
            f"{self.day},{self.matric(2)},maybe",  # bad status
            f"{self.semester.start_date - timedelta(days=1)},{self.matric(2)},P",  # outside semester
        ]
        data = self.post(csv_file(*rows)).data
        self.assertFalse(data["saved"])
        self.assertEqual(data["valid"], 1)
        problems = {e["row"]: str(e["errors"]) for e in data["errors"]}
        self.assertIn("Same student and class as row 2", problems[3])
        self.assertIn("isn't registered", problems[4])
        self.assertIn("future", problems[5])
        self.assertIn("Use a date like", problems[6])
        self.assertIn("present, late, excused or absent", problems[7])
        self.assertIn("outside", problems[8])
        self.assertFalse(AttendanceRecord.objects.exists())

    def test_closed_classes_need_a_correction(self):
        AttendanceSession.objects.create(
            offering=self.offering,
            date=self.day,
            start_time=self.offering.start_time,
            status=AttendanceSession.Status.CLOSED,
        )
        data = self.post(csv_file(f"{self.day},{self.matric(0)},present")).data
        self.assertIn("already closed", str(data["errors"]))

    def test_class_time_column(self):
        rows = [f"{self.day},{self.matric(0)},P,2pm", f"{self.day},{self.matric(1)},P,14:00"]
        self.post(csv_file(*rows, header="date,matric_number,status,start_time"), close="false")
        self.assertEqual(AttendanceSession.objects.get().start_time.hour, 14)

    def test_only_the_course_lecturer(self):
        other = client_for(make_staff(roles=["lecturer"]))
        self.assertEqual(other.post(self.url, {"file": csv_file()}, format="multipart").status_code, 403)
        self.assertEqual(client_for(self.students[0]).post(self.url, {}, format="multipart").status_code, 403)

    def test_template_is_the_class_list(self):
        response = self.client.get(
            f"/api/attendance/offerings/{self.offering.pk}/upload-template/?file=xlsx&date={self.day}"
        )
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual([c.value for c in sheet[1]], ["date", "matric_number", "student_name", "status"])
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet["A2"].value, str(self.day))

    def test_excel_upload(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Date", "Matric No.", "Attendance"])
        sheet.append([self.day, self.matric(0), "P"])
        buffer = BytesIO()
        book.save(buffer)
        file = SimpleUploadedFile("register.xlsx", buffer.getvalue())
        self.assertEqual(self.post(file, close="false").status_code, 201)
        self.assertEqual(AttendanceRecord.objects.get().status, Status.PRESENT)
