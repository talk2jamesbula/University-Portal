from datetime import timedelta
from io import BytesIO
from unittest import mock

from django.core.cache import cache
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.throttling import ScopedRateThrottle

from apps.academics.models import Enrollment
from apps.core.models import Notification
from apps.core.testing import (
    client_for,
    make_admin,
    make_course,
    make_department,
    make_offering,
    make_programme,
    make_semester,
    make_staff,
    make_student,
)

from .. import services
from ..models import AttendanceCorrection, AttendanceRecord, AttendanceSession

Status = AttendanceRecord.Status


class AttendanceTestCase(TestCase):
    def setUp(self):
        cache.clear()  # check-in throttle history
        today = timezone.localdate()
        self.semester = make_semester(start_date=today - timedelta(days=30), end_date=today + timedelta(days=120))
        self.department = make_department()
        self.lecturer = make_staff(roles=["lecturer"], department=self.department)
        self.course = make_course(self.department, code="CSC201")
        self.offering = make_offering(self.course, self.semester, lecturer=self.lecturer)
        programme = make_programme(self.department)
        self.alice = make_student(programme=programme, email="alice@example.com")
        self.bola = make_student(programme=programme)
        self.chidi = make_student(programme=programme)
        for student in (self.alice, self.bola, self.chidi):
            Enrollment.objects.create(student=student, offering=self.offering, status=Enrollment.Status.REGISTERED)
        self.outsider = make_student(programme=programme)  # not registered for the course

    def make_session(self, *, opened=True, **kw):
        session = AttendanceSession.objects.create(
            offering=self.offering,
            date=kw.pop("date", timezone.localdate()),
            start_time=kw.pop("start_time", self.offering.start_time),
            created_by=self.lecturer,
            **kw,
        )
        if opened:  # set directly: open_session only accepts today's classes
            session.status, session.opened_at = AttendanceSession.Status.OPEN, timezone.now()
            session.save()
        return session

    def check_in(self, student, session, *, device="", use_code=False):
        token, code, _ = services.current_codes(session)
        payload = {"code": code} if use_code else {"session": session.pk, "token": token}
        return client_for(student).post("/api/attendance/check-in/", {**payload, "device_id": device}, format="json")


class RotatingCodeTests(AttendanceTestCase):
    def test_codes_rotate_and_previous_window_is_accepted(self):
        session = self.make_session()
        with mock.patch("apps.attendance.services.time.time", return_value=1_000_000.0):
            token, code, _ = services.current_codes(session)
        with mock.patch("apps.attendance.services.time.time", return_value=1_000_020.0):  # next window
            self.assertTrue(services.code_is_valid(session, token=token))
            self.assertTrue(services.code_is_valid(session, code=code))
            self.assertNotEqual(services.current_codes(session)[0], token)
        with mock.patch("apps.attendance.services.time.time", return_value=1_000_045.0):  # two windows on
            self.assertFalse(services.code_is_valid(session, token=token))
            self.assertFalse(services.code_is_valid(session, code=code))

    def test_codes_are_specific_to_a_session(self):
        first = self.make_session()
        second = self.make_session(start_time=first.start_time.replace(hour=first.start_time.hour + 1))
        token, code, _ = services.current_codes(first)
        self.assertFalse(services.code_is_valid(second, token=token, code=code))


class CheckInTests(AttendanceTestCase):
    def test_qr_check_in_marks_present(self):
        session = self.make_session()
        response = self.check_in(self.alice, session, device="phone-a")
        self.assertEqual(response.status_code, 201, response.data)
        record = AttendanceRecord.objects.get(session=session, student=self.alice)
        self.assertEqual((record.status, record.method), (Status.PRESENT, AttendanceRecord.Method.QR))

    def test_typed_code_finds_the_open_class(self):
        session = self.make_session()
        response = self.check_in(self.bola, session, use_code=True)
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["course_code"], "CSC201")
        self.assertEqual(AttendanceRecord.objects.get(student=self.bola).method, AttendanceRecord.Method.CODE)

    def test_repeat_scan_is_idempotent(self):
        session = self.make_session()
        self.assertEqual(self.check_in(self.alice, session).status_code, 201)
        again = self.check_in(self.alice, session)
        self.assertEqual(again.status_code, 200)
        self.assertFalse(again.data["created"])
        self.assertIn("already marked present", again.data["message"])
        self.assertEqual(AttendanceRecord.objects.filter(session=session, student=self.alice).count(), 1)

    def test_one_device_for_several_students_is_flagged(self):
        session = self.make_session()
        self.check_in(self.alice, session, device="shared-phone")
        self.check_in(self.bola, session, device="shared-phone")
        self.assertFalse(AttendanceRecord.objects.get(student=self.alice).flagged)
        flagged = AttendanceRecord.objects.get(student=self.bola)
        self.assertTrue(flagged.flagged)
        self.assertIn("another student", flagged.flag_reason)

    def test_late_after_the_grace_period(self):
        session = self.make_session(late_after_minutes=10)
        session.opened_at = timezone.now() - timedelta(minutes=12)
        session.save()
        response = self.check_in(self.alice, session)
        self.assertEqual(response.data["status"], Status.LATE)
        self.assertIn("(late)", response.data["message"])

    def test_rejections(self):
        session = self.make_session(checkin_minutes=15)
        self.assertIn("aren't registered", str(self.check_in(self.outsider, session).data))
        bad = client_for(self.alice).post(
            "/api/attendance/check-in/", {"session": session.pk, "token": "0" * 16}, format="json"
        )
        self.assertIn("expired or is wrong", str(bad.data))

        session.opened_at = timezone.now() - timedelta(minutes=16)
        session.save()
        self.assertIn("Check-in for this class has closed", str(self.check_in(self.alice, session).data))

        scheduled = self.make_session(opened=False, date=timezone.localdate() - timedelta(days=1))
        self.assertIn("isn't open", str(self.check_in(self.alice, scheduled).data))
        self.assertFalse(AttendanceRecord.objects.exists())

    def test_staff_cannot_check_in(self):
        session = self.make_session()
        self.assertEqual(self.check_in(self.lecturer, session).status_code, 403)

    @mock.patch.object(ScopedRateThrottle, "THROTTLE_RATES", {"checkin": "3/minute"})
    def test_check_in_is_throttled(self):
        session = self.make_session()
        codes = [self.check_in(self.alice, session).status_code for _ in range(4)]
        self.assertEqual(codes, [201, 200, 200, 429])


class LecturerTests(AttendanceTestCase):
    def test_create_open_and_qr(self):
        client = client_for(self.lecturer)
        response = client.post(
            "/api/attendance/sessions/",
            {
                "offering": self.offering.pk,
                "date": str(timezone.localdate()),
                "start_time": "08:00",
                "topic": "Recursion",
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        pk = response.data["id"]
        self.assertEqual(
            client.get(f"/api/attendance/sessions/{pk}/qr/").data, {"accepting": False, "status": "scheduled"}
        )
        self.assertEqual(client.post(f"/api/attendance/sessions/{pk}/open/").data["status"], "open")
        qr = client.get(f"/api/attendance/sessions/{pk}/qr/").data
        self.assertTrue(qr["accepting"])
        self.assertRegex(qr["code"], r"^\d{6}$")
        self.assertIn(f"/portal/attend?s={pk}&t=", qr["url"])
        self.assertTrue(qr["svg"].startswith("<svg"))
        self.assertEqual((qr["checked_in"], qr["registered"]), (0, 3))

    def test_attendance_starts_only_on_the_class_day(self):
        tomorrow = self.make_session(opened=False, date=timezone.localdate() + timedelta(days=1))
        response = client_for(self.lecturer).post(f"/api/attendance/sessions/{tomorrow.pk}/open/")
        self.assertIn("day of the class", str(response.data))

    def test_past_class_can_be_marked_by_hand_then_closed(self):
        past = self.make_session(opened=False, date=timezone.localdate() - timedelta(days=2))
        services.mark(past, self.alice, Status.PRESENT, self.lecturer)
        services.close_session(past, self.lecturer)
        past.refresh_from_db()
        self.assertEqual(past.status, AttendanceSession.Status.CLOSED)
        self.assertEqual(past.records.filter(status=Status.ABSENT).count(), 2)

    def test_duplicate_slot_and_dates_outside_semester_are_rejected(self):
        client = client_for(self.lecturer)
        payload = {"offering": self.offering.pk, "date": str(timezone.localdate()), "start_time": "08:00"}
        self.assertEqual(client.post("/api/attendance/sessions/", payload, format="json").status_code, 201)
        self.assertEqual(client.post("/api/attendance/sessions/", payload, format="json").status_code, 400)
        late = {**payload, "date": str(self.semester.end_date + timedelta(days=1))}
        self.assertEqual(client.post("/api/attendance/sessions/", late, format="json").status_code, 400)

    def test_only_the_course_lecturer_runs_sessions(self):
        other = make_staff(roles=["lecturer"], department=self.department)
        client = client_for(other)
        payload = {"offering": self.offering.pk, "date": str(timezone.localdate()), "start_time": "08:00"}
        self.assertEqual(client.post("/api/attendance/sessions/", payload, format="json").status_code, 403)
        session = self.make_session()
        self.assertEqual(client.get(f"/api/attendance/sessions/{session.pk}/").status_code, 404)

    def test_manual_marking_and_close_marks_absentees(self):
        session = self.make_session()
        self.check_in(self.alice, session)
        client = client_for(self.lecturer)
        marked = client.post(
            f"/api/attendance/sessions/{session.pk}/mark/",
            {"student": self.bola.pk, "status": "excused"},
            format="json",
        )
        self.assertEqual(marked.status_code, 200, marked.data)
        outsider = client.post(
            f"/api/attendance/sessions/{session.pk}/mark/", {"student": self.outsider.pk, "status": "present"}
        )
        self.assertEqual(outsider.status_code, 400)

        self.assertEqual(client.post(f"/api/attendance/sessions/{session.pk}/close/").status_code, 200)
        statuses = dict(AttendanceRecord.objects.filter(session=session).values_list("student_id", "status"))
        self.assertEqual(
            statuses, {self.alice.pk: Status.PRESENT, self.bola.pk: Status.EXCUSED, self.chidi.pk: Status.ABSENT}
        )
        self.assertTrue(Notification.objects.filter(user=self.chidi, title__contains="marked absent").exists())
        self.assertTrue(Notification.objects.filter(user=self.chidi, title__contains="below 75%").exists())
        self.assertFalse(Notification.objects.filter(user=self.alice).exists())

        closed = client.post(
            f"/api/attendance/sessions/{session.pk}/mark/", {"student": self.chidi.pk, "status": "present"}
        )
        self.assertIn("correction request", str(closed.data))
        self.assertEqual(client.post(f"/api/attendance/sessions/{session.pk}/close/").status_code, 400)

    def test_roster_lists_every_registered_student(self):
        session = self.make_session()
        self.check_in(self.alice, session)
        data = client_for(self.lecturer).get(f"/api/attendance/sessions/{session.pk}/").data
        self.assertTrue(data["can_mark"])
        by_student = {row["student"]: row["record"] for row in data["roster"]}
        self.assertEqual(set(by_student), {self.alice.pk, self.bola.pk, self.chidi.pk})
        self.assertEqual(by_student[self.alice.pk]["status"], "present")
        self.assertIsNone(by_student[self.bola.pk])

    def test_only_unstarted_sessions_can_be_deleted(self):
        client = client_for(self.lecturer)
        scheduled = self.make_session(opened=False)
        started = self.make_session(start_time=scheduled.start_time.replace(hour=scheduled.start_time.hour + 1))
        self.assertEqual(client.delete(f"/api/attendance/sessions/{started.pk}/").status_code, 400)
        self.assertEqual(client.delete(f"/api/attendance/sessions/{scheduled.pk}/").status_code, 204)


class CorrectionTests(AttendanceTestCase):
    def setUp(self):
        super().setUp()
        self.session = self.make_session()
        services.close_session(self.session, self.lecturer)
        self.record = AttendanceRecord.objects.get(session=self.session, student=self.chidi)
        self.hod = make_staff(roles=[("hod", self.department)], department=self.department)

    def request(self, to_status="excused"):
        return client_for(self.lecturer).post(
            f"/api/attendance/records/{self.record.pk}/correction/",
            {"to_status": to_status, "reason": "Medical certificate presented."},
            format="json",
        )

    def test_request_and_approve(self):
        response = self.request()
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(Notification.objects.filter(user=self.hod, title__contains="correction").exists())
        self.assertEqual(self.request().status_code, 400)  # already pending

        queue = client_for(self.hod).get("/api/attendance/corrections/?status=pending").data["results"]
        self.assertEqual([c["id"] for c in queue], [response.data["id"]])
        decided = client_for(self.hod).post(
            f"/api/attendance/corrections/{response.data['id']}/decide/", {"approve": True, "note": "OK"}, format="json"
        )
        self.assertEqual(decided.data["status"], "approved")
        self.record.refresh_from_db()
        self.assertEqual((self.record.status, self.record.method), (Status.EXCUSED, AttendanceRecord.Method.CORRECTION))
        self.assertTrue(
            Notification.objects.filter(user=self.lecturer, title="Attendance correction approved").exists()
        )
        again = client_for(self.hod).post(
            f"/api/attendance/corrections/{response.data['id']}/decide/", {"approve": False}, format="json"
        )
        self.assertEqual(again.status_code, 400)

    def test_reject_leaves_the_record(self):
        pk = self.request().data["id"]
        client_for(self.hod).post(f"/api/attendance/corrections/{pk}/decide/", {"approve": False}, format="json")
        self.record.refresh_from_db()
        self.assertEqual(self.record.status, Status.ABSENT)

    def test_hod_of_another_department_cannot_approve(self):
        pk = self.request().data["id"]
        other_hod = make_staff(roles=[("hod", make_department())])
        response = client_for(other_hod).post(f"/api/attendance/corrections/{pk}/decide/", {"approve": True})
        self.assertEqual(response.status_code, 404)  # not even visible

    def test_requester_cannot_approve_their_own(self):
        acting = make_staff(roles=["lecturer", ("hod", self.department)], department=self.department)
        self.offering.lecturer = acting
        self.offering.save()
        correction = services.request_correction(self.record, Status.PRESENT, "Was in class", acting)
        response = client_for(acting).post(f"/api/attendance/corrections/{correction.pk}/decide/", {"approve": True})
        self.assertEqual(response.status_code, 403)

    def test_lecturer_cannot_decide_and_open_sessions_need_no_correction(self):
        pk = self.request().data["id"]
        response = client_for(self.lecturer).post(f"/api/attendance/corrections/{pk}/decide/", {"approve": True})
        self.assertEqual(response.status_code, 403)
        open_session = self.make_session(
            start_time=self.session.start_time.replace(hour=self.session.start_time.hour + 1)
        )
        record = services.mark(open_session, self.alice, Status.ABSENT, self.lecturer)
        self.assertEqual(AttendanceCorrection.objects.count(), 1)
        response = client_for(self.lecturer).post(
            f"/api/attendance/records/{record.pk}/correction/", {"to_status": "present", "reason": "Arrived late"}
        )
        self.assertIn("still open", str(response.data))


class StatsAndReportTests(AttendanceTestCase):
    def setUp(self):
        super().setUp()
        # Four closed sessions: Alice attends all, Bola 3 (one excused), Chidi 1.
        base = self.offering.start_time
        for i in range(4):
            session = self.make_session(date=timezone.localdate() - timedelta(days=i + 1), start_time=base)
            services.mark(session, self.alice, Status.PRESENT if i else Status.LATE, self.lecturer)
            if i < 3:
                services.mark(session, self.bola, Status.EXCUSED if i == 0 else Status.PRESENT, self.lecturer)
            if i == 0:
                services.mark(session, self.chidi, Status.PRESENT, self.lecturer)
            services.close_session(session, self.lecturer)
        self.make_session()  # an open session doesn't count yet
        self.registrar = make_staff(roles=["registrar"])

    def test_student_history_and_percentages(self):
        data = client_for(self.bola).get("/api/attendance/me/").data
        course = data["courses"][0]
        self.assertEqual((course["held"], course["attended"], course["absent"], course["percent"]), (4, 3, 1, 75.0))
        self.assertFalse(course["at_risk"])
        self.assertEqual(len(course["history"]), 4)
        self.assertEqual(data["overall_percent"], 75.0)
        chidi = client_for(self.chidi).get("/api/attendance/me/").data["courses"][0]
        self.assertEqual((chidi["percent"], chidi["at_risk"]), (25.0, True))

    def test_offering_stats(self):
        data = client_for(self.lecturer).get(f"/api/attendance/offerings/{self.offering.pk}/stats/").data
        self.assertEqual(data["summary"], {"sessions_held": 4, "average_percent": 66.7, "at_risk": 1, "students": 3})
        other = make_staff(roles=["lecturer"])
        self.assertEqual(client_for(other).get(f"/api/attendance/offerings/{self.offering.pk}/stats/").status_code, 403)

    def test_report_filters(self):
        client = client_for(self.registrar)
        data = client.get("/api/attendance/reports/").data
        self.assertEqual(data["summary"]["rows"], 3)
        self.assertEqual(data["summary"]["absences"], 4)
        self.assertEqual(client.get(f"/api/attendance/reports/?student={self.chidi.pk}").data["summary"]["rows"], 1)
        self.assertEqual(
            client.get("/api/attendance/reports/?at_risk=true").data["rows"][0]["student_id"], self.chidi.pk
        )
        self.assertEqual(
            client.get(f"/api/attendance/reports/?faculty={self.department.faculty_id + 999}").data["rows"], []
        )
        self.assertEqual(
            client.get(f"/api/attendance/reports/?department={self.department.pk}").data["summary"]["rows"], 3
        )
        self.assertEqual(
            client.get(f"/api/attendance/reports/?search={self.alice.university_id}").data["summary"]["rows"], 1
        )

    def test_report_is_scoped_and_permissioned(self):
        other_hod = make_staff(roles=[("hod", make_department())])
        self.assertEqual(client_for(other_hod).get("/api/attendance/reports/").data["rows"], [])
        own_hod = make_staff(roles=[("hod", self.department)])
        self.assertEqual(client_for(own_hod).get("/api/attendance/reports/").data["summary"]["rows"], 3)
        self.assertEqual(client_for(self.lecturer).get("/api/attendance/reports/").status_code, 403)
        self.assertEqual(client_for(self.alice).get("/api/attendance/reports/").status_code, 403)

    def test_exports(self):
        client = client_for(make_admin())
        csv = client.get("/api/attendance/reports/export/?file=csv")
        self.assertEqual(csv["Content-Type"], "text/csv; charset=utf-8")
        lines = csv.content.decode("utf-8-sig").strip().splitlines()
        self.assertTrue(lines[0].startswith("Matric no.,Student,Course"))
        self.assertEqual(len(lines), 4)

        xlsx = client.get("/api/attendance/reports/export/?file=xlsx&at_risk=true")
        sheet = load_workbook(BytesIO(xlsx.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["L2"].value, "Yes")

    @override_settings(ATTENDANCE_MIN_PERCENT=80)
    def test_minimum_is_configurable(self):
        data = client_for(self.bola).get("/api/attendance/me/").data
        self.assertTrue(data["courses"][0]["at_risk"])
