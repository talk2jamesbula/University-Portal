"""Reading uploaded CSV/Excel files and writing templates, for every bulk upload in the portal."""

import csv
from datetime import date, datetime
from io import BytesIO, StringIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.exceptions import ValidationError

DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def column_key(header, aliases=None):
    """'Matric No.' → 'matric_no' (then mapped through `aliases`)."""
    key = str(header or "").strip().lower().replace(" ", "_").replace(".", "").replace("-", "_")
    return (aliases or {}).get(key, key)


def read_rows(upload, *, required, aliases=None, max_rows=2000, noun="rows"):
    """[(spreadsheet row number, {column: value})] from a .csv or .xlsx upload; blank lines are skipped."""
    name = upload.name.lower()
    if name.endswith(".xlsx"):
        try:
            sheet = load_workbook(upload, read_only=True, data_only=True).active
        except Exception as exc:  # noqa: BLE001 - any unreadable workbook
            raise ValidationError({"file": "This Excel file can't be read."}) from exc
        rows = list(sheet.iter_rows(values_only=True))
    elif name.endswith(".csv"):
        try:
            text = upload.read().decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValidationError({"file": "Save the CSV file as UTF-8."}) from exc
        rows = list(csv.reader(StringIO(text)))
    else:
        raise ValidationError({"file": "Upload a .csv or .xlsx file."})
    if not rows:
        raise ValidationError({"file": "The file is empty."})
    header = [column_key(h, aliases) for h in rows[0]]
    missing = [c for c in required if c not in header]
    if missing:
        raise ValidationError({"file": f"Missing columns: {', '.join(missing)}. Download the template."})
    data = [
        (number, dict(zip(header, values, strict=False)))
        for number, values in enumerate(rows[1:], start=2)
        if any(v not in (None, "") for v in values)
    ]
    if len(data) > max_rows:
        raise ValidationError({"file": f"Upload at most {max_rows:,} {noun} at a time."})
    if not data:
        raise ValidationError({"file": f"The file has no {noun}."})
    return data


def text(value):
    """A cell as clean text (Excel gives whole numbers as floats)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_date(value):
    """A date from an Excel date cell or text such as 2026-10-05 or 05/10/2026 (day first). None if unreadable."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text(value), fmt).date()
        except ValueError:
            continue
    return None


def write_table(header, rows, file_format="csv", title="Sheet"):
    """A CSV (with BOM, so Excel reads UTF-8 names) or an Excel workbook with a styled header."""
    if file_format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = title[:31]
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12305A")
        for row in rows:
            sheet.append(row)
        for column in sheet.columns:
            width = max(len(str(c.value or "")) for c in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(10, width + 2), 40)
        sheet.freeze_panes = "A2"
        buffer = BytesIO()
        book.save(buffer)
        return buffer.getvalue()
    out = StringIO()
    csv.writer(out).writerows([header, *rows])
    return out.getvalue().encode("utf-8-sig")
