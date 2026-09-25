import csv
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db.models import Count, Q
from django.http import FileResponse, HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework import mixins, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import NotFound, PermissionDenied
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.permissions import BasePermission
from rest_framework.response import Response

from apps.academics.models import Enrollment, Semester
from apps.academics.serializers import SemesterSerializer
from apps.academics.services import current_semester, registration_summary, student_results
from apps.academics.views import serialize_results
from apps.accounts.models import LoginEvent, StudentProfile
from apps.accounts.views import save_avatar_upload
from apps.admissions.models import ApplicationDocument
from apps.attendance.services import student_attendance
from apps.core.models import AuditLog
from apps.core.services import audit
from apps.finance.serializers import ChargeSerializer, PaymentSerializer
from apps.finance.services import account_summary

from . import importer, services
from .models import StatusChange, StudentDocument
from .serializers import (
    ActivationSerializer,
    DocumentSerializer,
    DocumentUploadSerializer,
    StatusChangeRequestSerializer,
    StatusChangeSerializer,
    StudentDetailSerializer,
    StudentListSerializer,
    StudentWriteSerializer,
)

User = get_user_model()


class CanViewStudents(BasePermission):
    message = "You don't have access to student records."

    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and services.can_view(request.user))


# ?param=value → lookup on the User queryset
FILTERS = {
    "faculty": "student_profile__programme__department__faculty",
    "department": "student_profile__programme__department",
    "programme": "student_profile__programme",
    "level": "student_profile__level",
    "status": "student_profile__status",
    "gender": "student_profile__gender",
    "current_session": "student_profile__current_session",
    "entry_session": "student_profile__entry_session",
    "mode_of_entry": "student_profile__mode_of_entry",
    "state_of_origin": "student_profile__state_of_origin",
}
ORDERINGS = {
    "name": ["last_name", "first_name"],
    "-name": ["-last_name", "-first_name"],
    "matric": ["university_id"],
    "level": ["student_profile__level", "last_name"],
    "-level": ["-student_profile__level", "last_name"],
    "newest": ["-date_joined"],
}
EXPORT_COLUMNS = [
    ("Student ID", lambda u: u.student_profile.student_id),
    ("Matric no.", lambda u: u.university_id),
    ("Surname", lambda u: u.last_name),
    ("First name", lambda u: u.first_name),
    ("Gender", lambda u: u.student_profile.get_gender_display()),
    ("Date of birth", lambda u: u.student_profile.date_of_birth and u.student_profile.date_of_birth.isoformat()),
    ("Email", lambda u: u.email),
    ("Phone", lambda u: u.phone),
    ("Faculty", lambda u: u.student_profile.programme.department.faculty.name),
    ("Department", lambda u: u.student_profile.programme.department.name),
    ("Programme", lambda u: u.student_profile.programme.title),
    ("Level", lambda u: u.student_profile.level),
    ("Current session", lambda u: u.student_profile.current_session),
    ("Entry session", lambda u: u.student_profile.entry_session),
    ("Mode of entry", lambda u: u.student_profile.get_mode_of_entry_display()),
    ("Status", lambda u: u.student_profile.get_status_display()),
    ("Portal access", lambda u: "Active" if u.is_active else "Deactivated"),
    ("State of origin", lambda u: u.student_profile.state_of_origin),
    ("LGA", lambda u: u.student_profile.lga),
    ("Country", lambda u: u.student_profile.country),
    ("Address", lambda u: u.student_profile.home_address),
    ("Next of kin", lambda u: u.student_profile.next_of_kin_name),
    ("Next of kin phone", lambda u: u.student_profile.next_of_kin_phone),
]


def file_response(content, filename, content_type):
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


class StudentViewSet(mixins.ListModelMixin, mixins.RetrieveModelMixin, viewsets.GenericViewSet):
    """Student records for the Registry (all students) and HODs/Deans (their department or faculty)."""

    permission_classes = [CanViewStudents]
    serializer_class = StudentListSerializer
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    # --- Queryset -------------------------------------------------------------------------------

    def get_queryset(self):
        qs = services.visible_students(self.request.user).select_related(
            "student_profile__programme__department__faculty"
        )
        params = self.request.query_params
        for param, lookup in FILTERS.items():
            if params.get(param):
                qs = qs.filter(**{lookup: params[param]})
        if params.get("is_active") in ("true", "false"):
            qs = qs.filter(is_active=params["is_active"] == "true")
        if params.get("search", "").strip():
            qs = qs.filter(services.search_filter(params["search"]))
        return qs.order_by(*ORDERINGS.get(params.get("ordering"), ORDERINGS["name"]))

    def student(self):
        return self.get_object()

    def manager_only(self):
        services.require_manage(self.request.user)

    # --- List, create, read, update -------------------------------------------------------------

    def retrieve(self, request, *args, **kwargs):
        student = self.student()
        return Response(
            {
                **StudentDetailSerializer(student).data,
                "can_manage": services.can_manage(request.user),
                "can_view_fees": request.user.has_permission("finance.view", "students.manage"),
            }
        )

    def create(self, request):
        self.manager_only()
        serializer = StudentWriteSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        student, password = services.create_student(serializer.validated_data, request.user, request)
        return Response(
            {**StudentDetailSerializer(student).data, "temporary_password": password}, status=status.HTTP_201_CREATED
        )

    def update(self, request, pk=None):
        self.manager_only()
        student = self.student()
        serializer = StudentWriteSerializer(student, data=request.data)
        serializer.is_valid(raise_exception=True)
        services.update_student(student, serializer.validated_data, request.user, request)
        return Response(StudentDetailSerializer(User.objects.get(pk=student.pk)).data)

    partial_update = update

    @action(detail=False, methods=["get"])
    def summary(self, request):
        qs = self.get_queryset().order_by()
        by_status = dict(qs.values_list("student_profile__status").annotate(n=Count("id")))
        by_level = dict(qs.values_list("student_profile__level").annotate(n=Count("id")))
        return Response(
            {
                "total": qs.count(),
                "inactive": qs.filter(is_active=False).count(),
                "by_status": {s: by_status.get(s, 0) for s in StudentProfile.Status.values},
                "by_level": {str(k): v for k, v in sorted(by_level.items())},
                "can_manage": services.can_manage(request.user),
            }
        )

    # --- Status, access, photo, password --------------------------------------------------------

    @action(detail=True, methods=["post"])
    def status(self, request, pk=None):
        self.manager_only()
        serializer = StatusChangeRequestSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        services.change_status(
            self.student(), data["status"], data["reason"].strip(), request.user, request, data.get("effective_date")
        )
        return self.retrieve(request)

    @action(detail=True, methods=["post"])
    def activation(self, request, pk=None):
        self.manager_only()
        serializer = ActivationSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        services.set_active(self.student(), data["active"], data["reason"].strip(), request.user, request)
        return self.retrieve(request)

    @action(detail=True, methods=["post", "delete"])
    def photo(self, request, pk=None):
        self.manager_only()
        student = self.student()
        if request.method == "DELETE":
            student.set_avatar(None)
        else:
            save_avatar_upload(student, request)
        audit(
            request,
            "students.photo",
            student,
            f"{'Removed' if request.method == 'DELETE' else 'Changed'} the passport photograph "
            f"of {student.university_id}",
        )
        return self.retrieve(request)

    @action(detail=True, methods=["post"], url_path="reset-password")
    def reset_password(self, request, pk=None):
        self.manager_only()
        return Response({"temporary_password": services.reset_password(self.student(), request.user, request)})

    # --- Profile tabs ---------------------------------------------------------------------------

    @action(detail=True, methods=["get"], url_path="academic-records")
    def academic_records(self, request, pk=None):
        student = self.student()
        results = student_results(student)
        semesters = (
            Semester.objects.filter(
                offerings__enrollments__student=student, offerings__enrollments__status=Enrollment.Status.REGISTERED
            )
            .distinct()
            .order_by("-start_date")
        )
        registrations = []
        for semester in semesters:
            summary = registration_summary(student, semester)
            registrations.append(
                {
                    "semester": SemesterSerializer(semester).data,
                    "units": summary["units"],
                    "courses": len(summary["registered"]),
                }
            )
        return Response(
            {
                "cgpa": results["cgpa"],
                "standing": results["standing"],
                "degree_class": results["degree_class"],
                "units_taken": results["units_taken"],
                "units_passed": results["units_passed"],
                "semesters": [
                    {
                        "semester": SemesterSerializer(r["semester"]).data,
                        "gpa": r["gpa"],
                        "units_taken": r["units_taken"],
                        "units_passed": r["units_passed"],
                    }
                    for r in results["semesters"]
                ],
                "registrations": registrations,
                "status_history": StatusChangeSerializer(
                    StatusChange.objects.filter(student=student).select_related("changed_by"), many=True
                ).data,
            }
        )

    @action(detail=True, methods=["get"])
    def courses(self, request, pk=None):
        student = self.student()
        enrollments = (
            Enrollment.objects.filter(student=student, status=Enrollment.Status.REGISTERED)
            .select_related("offering__course", "offering__semester", "offering__lecturer")
            .order_by("-offering__semester__start_date", "offering__course__code")
        )
        groups = {}
        for e in enrollments:
            semester = e.offering.semester
            group = groups.setdefault(
                semester.pk, {"semester": SemesterSerializer(semester).data, "units": 0, "courses": []}
            )
            group["units"] += e.offering.course.units
            group["courses"].append(
                {
                    "offering": e.offering_id,
                    "code": e.offering.course.code,
                    "title": e.offering.course.title,
                    "units": e.offering.course.units,
                    "lecturer": e.offering.lecturer.get_full_name() if e.offering.lecturer else None,
                    "is_carryover": e.is_carryover,
                    "result_status": e.result_status,
                    "grade": e.grade if e.result_status == Enrollment.ResultStatus.PUBLISHED else None,
                }
            )
        semester = current_semester()
        return Response({"current_semester": semester.pk if semester else None, "semesters": list(groups.values())})

    @action(detail=True, methods=["get"])
    def attendance(self, request, pk=None):
        return Response(student_attendance(self.student()))

    @action(detail=True, methods=["get"])
    def results(self, request, pk=None):
        return Response(serialize_results(student_results(self.student())))

    @action(detail=True, methods=["get"])
    def fees(self, request, pk=None):
        if not request.user.has_permission("finance.view", "students.manage"):
            raise PermissionDenied("You don't have access to student fees.")
        summary = account_summary(self.student())
        return Response(
            {
                **{k: summary[k] for k in ("total_charged", "total_paid", "balance", "overdue", "next_due_date")},
                "charges": ChargeSerializer(summary["charges"], many=True).data,
                "payments": PaymentSerializer(summary["payments"], many=True).data,
            }
        )

    @action(detail=True, methods=["get", "post"])
    def documents(self, request, pk=None):
        student = self.student()
        if request.method == "POST":
            self.manager_only()
            serializer = DocumentUploadSerializer(data=request.data)
            serializer.is_valid(raise_exception=True)
            upload = serializer.validated_data["file"]
            document = StudentDocument.objects.create(
                student=student,
                kind=serializer.validated_data["kind"],
                title=serializer.validated_data["title"],
                file=upload,
                original_filename=upload.name[:200],
                content_type=serializer.content_type,
                size=upload.size,
                uploaded_by=request.user,
            )
            audit(
                request,
                "students.document_upload",
                student,
                f"Added {document.get_kind_display()} to {student.university_id}",
            )
            return Response(DocumentSerializer(document).data, status=status.HTTP_201_CREATED)
        documents = StudentDocument.objects.filter(student=student).select_related("uploaded_by")
        # Documents from the student's online application, if they applied through the portal.
        admission = (
            ApplicationDocument.objects.filter(
                application__applicant__email__iexact=student.email, application__status__in=["admitted", "accepted"]
            )
            if student.email
            else ApplicationDocument.objects.none()
        )
        return Response(
            {
                "documents": DocumentSerializer(documents, many=True).data,
                "admission_documents": [
                    {
                        "id": d.id,
                        "kind_label": d.get_kind_display(),
                        "original_filename": d.original_filename,
                        "status_label": d.get_status_display(),
                        "uploaded_at": d.uploaded_at,
                    }
                    for d in admission
                ],
            }
        )

    @action(detail=True, methods=["get"])
    def activity(self, request, pk=None):
        """What happened to this record (audit log, status changes) and what the student did (sign-ins, actions)."""
        student = self.student()
        events = [
            {
                "at": log.created_at,
                "kind": "record" if log.actor_id != student.pk else "student",
                "action": log.action,
                "summary": log.summary,
                "actor": log.actor.get_full_name() if log.actor else "System",
            }
            for log in AuditLog.objects.filter(
                Q(target_type="accounts.user", target_id=str(student.pk)) | Q(actor=student)
            ).select_related("actor")[:150]
        ]
        events += [
            {
                "at": e.created_at,
                "kind": "login",
                "action": "login" if e.successful else "login_failed",
                "summary": f"{'Signed in' if e.successful else 'Failed sign-in'} "
                f"from {e.ip_address or 'unknown address'}",
                "actor": student.get_full_name(),
            }
            for e in LoginEvent.objects.filter(user=student)[:50]
        ]
        events.sort(key=lambda e: e["at"], reverse=True)
        return Response(events[:150])

    # --- Import and export ----------------------------------------------------------------------

    @action(detail=False, methods=["get"])
    def export(self, request):
        students = list(self.get_queryset())
        audit(request, "students.export", None, f"Exported {len(students)} student records")
        header = [label for label, _ in EXPORT_COLUMNS]
        rows = [[get(u) for _, get in EXPORT_COLUMNS] for u in students]
        stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
        if request.query_params.get("file") == "xlsx":
            book = Workbook()
            sheet = book.active
            sheet.title = "Students"
            sheet.append(header)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="12305A")
            for row in rows:
                sheet.append(row)
            for column in sheet.columns:
                width = max(len(str(c.value or "")) for c in column)
                sheet.column_dimensions[column[0].column_letter].width = min(max(10, width + 2), 40)
            sheet.freeze_panes = "C2"
            buffer = BytesIO()
            book.save(buffer)
            return file_response(
                buffer.getvalue(),
                f"students-{stamp}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="students-{stamp}.csv"'
        response.write("﻿")
        writer = csv.writer(response)
        writer.writerow(header)
        writer.writerows(["" if v is None else v for v in row] for row in rows)
        return response

    @action(detail=False, methods=["get"], url_path="import-template")
    def import_template(self, request):
        self.manager_only()
        if request.query_params.get("file") == "xlsx":
            return file_response(
                importer.template("xlsx"),
                "student-import-template.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        return file_response(importer.template("csv"), "student-import-template.csv", "text/csv; charset=utf-8")

    @action(detail=False, methods=["post"], url_path="import")
    def import_students(self, request):
        self.manager_only()
        upload = request.FILES.get("file")
        if not upload:
            return Response({"file": ["Choose a CSV or Excel file."]}, status=status.HTTP_400_BAD_REQUEST)
        dry_run = str(request.data.get("dry_run", "")).lower() in ("1", "true", "yes")
        result = importer.import_students(upload, request.user, request, dry_run=dry_run)
        code = status.HTTP_201_CREATED if result["created"] else status.HTTP_200_OK
        return Response(result, status=code)


class DocumentFileViewSet(viewsets.GenericViewSet):
    """Download or delete one student document (the Registry, or staff who can see that student)."""

    permission_classes = [CanViewStudents]

    def get_document(self, pk):
        document = get_object_or_404(StudentDocument.objects.select_related("student"), pk=pk)
        if not services.visible_students(self.request.user).filter(pk=document.student_id).exists():
            raise NotFound()
        return document

    def retrieve(self, request, pk=None):
        document = self.get_document(pk)
        return FileResponse(
            document.file.open("rb"), content_type=document.content_type, filename=document.original_filename
        )

    def destroy(self, request, pk=None):
        services.require_manage(request.user)
        document = self.get_document(pk)
        name = document.file.name
        audit(
            request,
            "students.document_delete",
            document.student,
            f"Removed {document.get_kind_display()} from {document.student.university_id}",
        )
        document.delete()
        document.file.storage.delete(name)
        return Response(status=status.HTTP_204_NO_CONTENT)
