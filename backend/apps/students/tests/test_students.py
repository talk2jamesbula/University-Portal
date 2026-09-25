import shutil
import tempfile
from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import StudentProfile
from apps.core.models import AuditLog, Notification
from apps.core.testing import client_for, make_department, make_faculty, make_programme, make_staff, make_student

from ..models import StatusChange, StudentDocument

User = get_user_model()
MEDIA = tempfile.mkdtemp()
PDF = b"%PDF-1.4\n%fake\n"
URL = "/api/students/"


def payload(**overrides):
    return {
        "first_name": "Adaeze",
        "last_name": "Okafor",
        "email": "adaeze@example.com",
        "phone": "08031234567",
        "gender": "female",
        "date_of_birth": "2007-05-14",
        "level": 100,
        "entry_session": "2026/2027",
        "current_session": "2026/2027",
        "mode_of_entry": "utme",
        "state_of_origin": "Enugu",
        "country": "Nigeria",
        "home_address": "12 Okpara Avenue",
        "next_of_kin_name": "Mrs. Ngozi Okafor",
        "next_of_kin_phone": "08037654321",
        **overrides,
    }


@override_settings(MEDIA_ROOT=MEDIA, EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class StudentsTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.science = make_faculty(name="Science")
        cls.csc_dept = make_department(cls.science, name="Computer Science")
        cls.mth_dept = make_department(cls.science, name="Mathematics")
        cls.arts_dept = make_department(make_faculty(name="Arts"), name="English")
        cls.csc = make_programme(cls.csc_dept, code="CSC", name="Computer Science")
        cls.mth = make_programme(cls.mth_dept, code="MTH", name="Mathematics")
        cls.eng = make_programme(cls.arts_dept, code="ENG", name="English")
        cls.registrar = make_staff("registrar", roles=["registrar"])
        cls.hod = make_staff("hod", roles=[("hod", cls.csc_dept)], department=cls.csc_dept)
        cls.dean = make_staff("dean", roles=[("dean", cls.science)])
        cls.lecturer = make_staff("lect", roles=["lecturer"])
        cls.ada = make_student(
            "ada", programme=cls.csc, level=200, email="ada@example.com", first_name="Ada", last_name="Obi"
        )
        cls.bayo = make_student(
            "bayo", programme=cls.mth, level=300, email="bayo@example.com", first_name="Bayo", last_name="Ade"
        )
        cls.chi = make_student(
            "chi", programme=cls.eng, level=100, email="chi@example.com", first_name="Chi", last_name="Eze"
        )

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        self.reg = client_for(self.registrar)

    def create(self, **overrides):
        return self.reg.post(URL, payload(programme=self.csc.pk, **overrides), format="json")


class AccessTests(StudentsTestCase):
    def ids(self, user):
        return {r["id"] for r in client_for(user).get(URL).data["results"]}

    def test_scope(self):
        self.assertEqual(self.ids(self.registrar), {self.ada.pk, self.bayo.pk, self.chi.pk})
        self.assertEqual(self.ids(self.hod), {self.ada.pk})  # own department
        self.assertEqual(self.ids(self.dean), {self.ada.pk, self.bayo.pk})  # own faculty
        self.assertEqual(client_for(self.hod).get(f"{URL}{self.chi.pk}/").status_code, 404)
        self.assertEqual(client_for(self.lecturer).get(URL).status_code, 403)
        self.assertEqual(client_for(self.ada).get(URL).status_code, 403)
        self.assertEqual(APIClient().get(URL).status_code, 401)

    def test_viewers_cannot_change_records(self):
        hod = client_for(self.hod)
        self.assertEqual(hod.post(URL, payload(programme=self.csc.pk), format="json").status_code, 403)
        self.assertEqual(
            hod.put(f"{URL}{self.ada.pk}/", payload(programme=self.csc.pk), format="json").status_code, 403
        )
        status_change = {"status": "suspended", "reason": "Misconduct"}
        self.assertEqual(hod.post(f"{URL}{self.ada.pk}/status/", status_change, format="json").status_code, 403)
        self.assertFalse(hod.get(f"{URL}{self.ada.pk}/").data["can_manage"])

    def test_fees_need_finance_or_registry(self):
        self.assertEqual(client_for(self.hod).get(f"{URL}{self.ada.pk}/fees/").status_code, 403)
        self.assertEqual(self.reg.get(f"{URL}{self.ada.pk}/fees/").status_code, 200)
        bursar = make_staff("bursar", roles=["bursar"])
        self.assertEqual(client_for(bursar).get(URL).status_code, 403)  # finance alone isn't student records


class CreateAndEditTests(StudentsTestCase):
    def test_create_generates_numbers_and_password(self):
        response = self.create()
        self.assertEqual(response.status_code, 201, response.data)
        user = User.objects.get(email="adaeze@example.com")
        self.assertEqual(user.university_id, "BU/26/CSC/0001")
        self.assertEqual(user.username, "bu26csc0001")
        self.assertEqual(user.department, self.csc_dept)
        self.assertTrue(user.student_profile.student_id.startswith("STU"))
        self.assertTrue(user.check_password(response.data["temporary_password"]))
        self.assertTrue(AuditLog.objects.filter(action="students.create", target_id=str(user.pk)).exists())
        self.assertTrue(Notification.objects.filter(user=user, title="Welcome to the student portal").exists())
        second = self.create(email="b@example.com")
        self.assertEqual(second.data["matric_number"], "BU/26/CSC/0002")
        login = APIClient().post(
            "/api/auth/token/", {"username": "bu26csc0001", "password": response.data["temporary_password"]}
        )
        self.assertEqual(login.status_code, 200)

    def test_validation(self):
        response = self.create(
            email="ada@example.com",
            phone="12",
            level=1000,
            current_session="2025/2027",
            date_of_birth="2024-01-01",
            jamb_reg_number="123",
        )
        self.assertEqual(
            set(response.data),
            {"email", "phone", "current_session", "date_of_birth", "jamb_reg_number"},
        )
        self.assertIn("level", self.create(level=150).data)
        self.assertIn(
            "before the entry session", str(self.create(entry_session="2026/2027", current_session="2025/2026").data)
        )
        self.mth.is_active = False
        self.mth.save()
        self.assertIn("not admitting", str(self.reg.post(URL, payload(programme=self.mth.pk), format="json").data))

    def test_edit_moves_department_and_is_audited(self):
        data = payload(programme=self.eng.pk, email="ada@example.com", first_name="Ada", last_name="Obi", level=300)
        response = self.reg.put(f"{URL}{self.ada.pk}/", data, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.ada.refresh_from_db()
        self.assertEqual(self.ada.department, self.arts_dept)
        self.assertEqual(self.ada.student_profile.level, 300)
        log = AuditLog.objects.get(action="students.update")
        self.assertIn("level: 200 → 300", log.summary)
        self.assertIn("programme:", log.summary)
        # The HOD of Computer Science no longer sees her.
        self.assertEqual(client_for(self.hod).get(f"{URL}{self.ada.pk}/").status_code, 404)


class StatusAndAccessTests(StudentsTestCase):
    def test_status_change_is_recorded_and_notified(self):
        url = f"{URL}{self.ada.pk}/status/"
        self.assertEqual(self.reg.post(url, {"status": "expelled", "reason": "x"}, format="json").status_code, 400)
        response = self.reg.post(
            url, {"status": "expelled", "reason": "Examination malpractice (SDC ref 12)"}, format="json"
        )
        self.assertEqual(response.data["profile"]["status"], "expelled")
        change = StatusChange.objects.get()
        self.assertEqual(
            (change.from_status, change.to_status, change.changed_by), ("active", "expelled", self.registrar)
        )
        self.assertTrue(Notification.objects.filter(user=self.ada, title__contains="Expelled").exists())
        self.assertTrue(AuditLog.objects.filter(action="students.status").exists())
        self.assertIn(
            "already", str(self.reg.post(url, {"status": "expelled", "reason": "Again, again"}, format="json").data)
        )
        history = self.reg.get(f"{URL}{self.ada.pk}/academic-records/").data["status_history"]
        self.assertEqual(history[0]["to_label"], "Expelled")

    def test_all_statuses_supported(self):
        self.assertTrue(
            {"active", "graduated", "suspended", "withdrawn", "deferred", "expelled", "completed"}
            <= set(StudentProfile.Status.values)
        )

    def test_deactivate_blocks_sign_in(self):
        self.ada.set_password("Pass-1234!")
        self.ada.save()
        self.reg.post(
            f"{URL}{self.ada.pk}/activation/", {"active": False, "reason": "Left the university"}, format="json"
        )
        self.ada.refresh_from_db()
        self.assertFalse(self.ada.is_active)
        login = APIClient().post("/api/auth/token/", {"username": "ada", "password": "Pass-1234!"})
        self.assertEqual(login.status_code, 401)
        self.assertEqual(self.reg.get(f"{URL}?is_active=false").data["count"], 1)
        self.reg.post(f"{URL}{self.ada.pk}/activation/", {"active": True}, format="json")
        self.assertTrue(AuditLog.objects.filter(action="students.activate").exists())

    def test_reset_password(self):
        password = self.reg.post(f"{URL}{self.ada.pk}/reset-password/").data["temporary_password"]
        self.ada.refresh_from_db()
        self.assertTrue(self.ada.check_password(password))


class SearchFilterTests(StudentsTestCase):
    def count(self, query):
        return self.reg.get(f"{URL}?{query}").data["count"]

    def test_search_and_filters(self):
        self.assertEqual(self.count("search=ada obi"), 1)
        self.assertEqual(self.count("search=obi ada"), 1)
        self.assertEqual(self.count(f"search={self.bayo.university_id}"), 1)
        self.assertEqual(self.count(f"search={self.chi.student_profile.student_id}"), 1)
        self.assertEqual(self.count(f"faculty={self.science.pk}"), 2)
        self.assertEqual(self.count(f"department={self.mth_dept.pk}"), 1)
        self.assertEqual(self.count(f"programme={self.eng.pk}&level=100"), 1)
        self.assertEqual(self.count("level=300"), 1)
        self.assertEqual(self.count("status=graduated"), 0)
        names = [r["last_name"] for r in self.reg.get(f"{URL}?ordering=-level").data["results"]]
        self.assertEqual(names, ["Ade", "Obi", "Eze"])

    def test_pagination_and_summary(self):
        self.assertEqual(set(self.reg.get(URL).data), {"count", "next", "previous", "results"})
        summary = self.reg.get(f"{URL}summary/").data
        self.assertEqual((summary["total"], summary["by_status"]["active"], summary["by_level"]["300"]), (3, 3, 1))
        self.assertEqual(client_for(self.hod).get(f"{URL}summary/").data["total"], 1)


class ImportExportTests(StudentsTestCase):
    def upload(self, rows, name="students.csv", dry_run=False):
        header = "first_name,last_name,email,gender,programme_code,level,entry_session,current_session,date_of_birth\n"
        body = (header + "\n".join(rows)).encode()
        file = SimpleUploadedFile(name, body, content_type="text/csv")
        return self.reg.post(f"{URL}import/", {"file": file, "dry_run": "true" if dry_run else ""}, format="multipart")

    def test_valid_import_creates_everyone(self):
        rows = [
            "Tolu,Bello,tolu@example.com,female,CSC,100,2026/2027,2026/2027,14/05/2007",
            "Musa,Garba,musa@example.com,Male,mth,100,2026/2027,2026/2027,2007-01-02",
        ]
        preview = self.upload(rows, dry_run=True).data
        self.assertEqual((preview["valid"], preview["created"]), (2, []))
        self.assertFalse(User.objects.filter(email="tolu@example.com").exists())
        result = self.upload(rows)
        self.assertEqual(result.status_code, 201, result.data)
        self.assertEqual([c["matric_number"] for c in result.data["created"]], ["BU/26/CSC/0001", "BU/26/MTH/0001"])
        self.assertEqual(
            User.objects.get(email="musa@example.com").student_profile.date_of_birth.isoformat(), "2007-01-02"
        )
        self.assertTrue(AuditLog.objects.filter(action="students.import").exists())

    def test_any_invalid_row_stops_the_import(self):
        rows = [
            "Tolu,Bello,tolu@example.com,female,CSC,100,2026/2027,2026/2027,",
            "Dup,Row,tolu@example.com,female,CSC,100,2026/2027,2026/2027,",
            "Bad,Prog,x@example.com,female,XYZ,100,2026/2027,2026/2027,",
            "Taken,Email,ada@example.com,female,CSC,100,2026/2027,2026/2027,",
        ]
        data = self.upload(rows).data
        self.assertEqual(data["valid"], 1)
        errors = {e["row"]: e["errors"] for e in data["errors"]}
        self.assertIn("Same as row 2", str(errors[3]))
        self.assertIn("XYZ", str(errors[4]))
        self.assertIn("email", errors[5])
        self.assertFalse(User.objects.filter(email="tolu@example.com").exists())

    def test_rejects_bad_files(self):
        self.assertIn(
            "Missing columns",
            str(
                self.reg.post(
                    f"{URL}import/", {"file": SimpleUploadedFile("s.csv", b"name,email\nA,b@c.com")}, format="multipart"
                ).data
            ),
        )
        self.assertIn(
            ".csv or .xlsx",
            str(self.reg.post(f"{URL}import/", {"file": SimpleUploadedFile("s.pdf", PDF)}, format="multipart").data),
        )
        self.assertEqual(client_for(self.hod).post(f"{URL}import/", {}, format="multipart").status_code, 403)

    def test_excel_import_and_template(self):
        template = load_workbook(BytesIO(self.reg.get(f"{URL}import-template/?file=xlsx").content)).active
        self.assertEqual(template["A1"].value, "first_name")
        book = Workbook()
        sheet = book.active
        sheet.append(
            ["First Name", "Last Name", "Email", "Gender", "Programme", "Level", "Entry Session", "Current Session"]
        )
        sheet.append(["Zara", "Umar", "zara@example.com", "female", "ENG", 200, "2025/2026", "2026/2027"])
        buffer = BytesIO()
        book.save(buffer)
        file = SimpleUploadedFile("students.xlsx", buffer.getvalue())
        response = self.reg.post(f"{URL}import/", {"file": file}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["created"][0]["matric_number"], "BU/25/ENG/0001")

    def test_export_respects_filters_and_scope(self):
        csv = self.reg.get(f"{URL}export/?file=csv&faculty={self.science.pk}")
        lines = csv.content.decode("utf-8-sig").strip().splitlines()
        self.assertTrue(lines[0].startswith("Student ID,Matric no.,Surname"))
        self.assertEqual(len(lines), 3)
        xlsx = client_for(self.hod).get(f"{URL}export/?file=xlsx")
        sheet = load_workbook(BytesIO(xlsx.content)).active
        self.assertEqual(sheet.max_row, 2)  # the HOD only exports their department
        self.assertTrue(AuditLog.objects.filter(action="students.export").exists())


class ProfileTabTests(StudentsTestCase):
    def test_tabs_load(self):
        for tab in ("academic-records", "courses", "attendance", "results", "fees", "documents", "activity"):
            response = self.reg.get(f"{URL}{self.ada.pk}/{tab}/")
            self.assertEqual(response.status_code, 200, tab)

    def test_documents_are_private(self):
        file = SimpleUploadedFile("letter.pdf", PDF)
        created = self.reg.post(
            f"{URL}{self.ada.pk}/documents/",
            {"kind": "letter", "title": "Suspension letter", "file": file},
            format="multipart",
        )
        self.assertEqual(created.status_code, 201, created.data)
        url = f"{URL}documents/{created.data['id']}/"
        self.assertEqual(client_for(self.hod).get(url).status_code, 200)
        self.assertEqual(client_for(self.dean).get(url).status_code, 200)
        other_hod = make_staff("hod2", roles=[("hod", self.arts_dept)])
        self.assertEqual(client_for(other_hod).get(url).status_code, 404)
        self.assertEqual(client_for(self.hod).delete(url).status_code, 403)
        self.assertEqual(self.reg.delete(url).status_code, 204)
        self.assertFalse(StudentDocument.objects.exists())
        fake = SimpleUploadedFile("x.pdf", b"not a pdf")
        self.assertEqual(
            self.reg.post(
                f"{URL}{self.ada.pk}/documents/", {"kind": "other", "file": fake}, format="multipart"
            ).status_code,
            400,
        )

    def test_activity_lists_record_changes_and_sign_ins(self):
        self.reg.post(
            f"{URL}{self.ada.pk}/status/", {"status": "suspended", "reason": "Unpaid fees review"}, format="json"
        )
        self.ada.set_password("Pass-1234!")
        self.ada.save()
        APIClient().post("/api/auth/token/", {"username": "ada", "password": "Pass-1234!"})
        kinds = {e["kind"] for e in self.reg.get(f"{URL}{self.ada.pk}/activity/").data}
        self.assertTrue({"record", "login"} <= kinds)


class AutomaticNumberTests(StudentsTestCase):
    def test_matric_numbers_cannot_be_typed_in(self):
        response = self.create(matric_number="BU/99/CSC/9999")
        self.assertEqual(response.data["matric_number"], "BU/26/CSC/0001")
        # Editing doesn't change it either.
        data = payload(programme=self.csc.pk, email="adaeze@example.com", matric_number="BU/26/CSC/0500")
        self.reg.put(f"{URL}{response.data['id']}/", data, format="json")
        self.assertEqual(User.objects.get(pk=response.data["id"]).university_id, "BU/26/CSC/0001")

    def test_numbers_follow_on_per_programme_and_year(self):
        numbers = [self.create(email=f"s{i}@example.com").data["matric_number"] for i in range(2)]
        numbers.append(
            self.reg.post(URL, payload(programme=self.mth.pk, email="m@example.com"), format="json").data[
                "matric_number"
            ]
        )
        numbers.append(
            self.create(email="o@example.com", entry_session="2025/2026", current_session="2026/2027").data[
                "matric_number"
            ]
        )
        self.assertEqual(numbers, ["BU/26/CSC/0001", "BU/26/CSC/0002", "BU/26/MTH/0001", "BU/25/CSC/0001"])

    def test_a_number_taken_at_the_same_moment_is_skipped(self):
        from apps.accounts import numbering

        first = self.create(email="a@example.com").data["matric_number"]  # BU/26/CSC/0001
        calls = []
        real = numbering._next

        def stale(prefix, first=1):  # the first attempt sees stale data and picks a number already in use
            calls.append(prefix)
            return 1 if len(calls) == 1 else real(prefix, first)

        with mock.patch("apps.accounts.numbering._next", side_effect=stale):
            second = self.create(email="b@example.com")
        self.assertEqual(second.status_code, 201, second.data)
        self.assertEqual((first, second.data["matric_number"]), ("BU/26/CSC/0001", "BU/26/CSC/0002"))

    def test_student_profile_created_elsewhere_still_gets_numbers(self):
        user = User.objects.create_user("new", password="x", role=User.Role.STUDENT)
        profile = StudentProfile.objects.create(user=user, programme=self.eng, entry_session="2024/2025")
        user.refresh_from_db()
        self.assertEqual(user.university_id, "BU/24/ENG/0001")
        self.assertTrue(profile.student_id.startswith("STU"))

    def test_staff_and_admin_numbers(self):
        a = User.objects.create_user("s1", password="x", role=User.Role.STAFF)
        b = User.objects.create_user("s2", password="x", role=User.Role.STAFF)
        admin = User.objects.create_user("a1", password="x", role=User.Role.ADMIN)
        self.assertEqual((a.university_id, b.university_id), ("SP/1001", "SP/1002"))
        self.assertTrue(admin.university_id.startswith("SA/"))
        # Through the users API the number can't be set by hand either.
        client = client_for(make_staff("reg2", roles=["registrar"]))
        response = client.post(
            "/api/users/",
            {"username": "s3", "first_name": "T", "last_name": "O", "role": "staff", "university_id": "X/1"},
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(response.data["university_id"].startswith("SP/"))

    def test_import_assigns_numbers_and_ignores_a_matric_column(self):
        body = (
            b"first_name,last_name,email,gender,programme_code,level,entry_session,current_session,matric_number\n"
            b"Tolu,Bello,tolu@example.com,female,CSC,100,2026/2027,2026/2027,BU/26/CSC/7777\n"
        )
        file = SimpleUploadedFile("s.csv", body)
        created = self.reg.post(f"{URL}import/", {"file": file}, format="multipart").data["created"]
        self.assertEqual(created[0]["matric_number"], "BU/26/CSC/0001")
