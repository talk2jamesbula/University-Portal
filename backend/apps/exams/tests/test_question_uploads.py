import csv
from datetime import timedelta
from io import BytesIO, StringIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.core.models import AuditLog
from apps.core.testing import client_for, make_staff

from ..models import Exam, Question
from .test_exams import ExamTestCase

HEADER = ["question", "option_a", "option_b", "option_c", "option_d", "answer", "marks"]


def csv_file(rows, header=HEADER, name="questions.csv"):
    out = StringIO()
    csv.writer(out).writerows([header, *rows])
    return SimpleUploadedFile(name, out.getvalue().encode("utf-8"), content_type="text/csv")


def xlsx_file(rows, header=HEADER):
    book = Workbook()
    sheet = book.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return SimpleUploadedFile("questions.xlsx", buffer.getvalue())


GOOD = [
    ["What is 2 + 2?", "3", "4", "5", "", "B", "1"],
    ["The earth orbits the sun.", "True", "False", "", "", "True", ""],
    ["Pick the prime number.", "4", "6", "7", "9", "c", "2"],
]


class QuestionUploadTests(ExamTestCase):
    def setUp(self):
        super().setUp()
        self.exam = self.make_exam(venues=[self.lab], mode=Exam.Mode.CBT)
        self.client_ = client_for(self.lecturer)
        self.url = f"/api/exams/timetable/{self.exam.pk}/questions/upload/"

    def post(self, file, **extra):
        return self.client_.post(self.url, {"file": file, **extra}, format="multipart")

    def test_dry_run_checks_without_saving(self):
        response = self.post(csv_file(GOOD), dry_run="true")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual((response.data["valid"], response.data["marks"], response.data["saved"]), (3, 4, False))
        self.assertEqual(response.data["preview"][1]["correct"], 0)  # "True" matched by text
        self.assertFalse(Question.objects.exists())

    def test_import_saves_questions_options_and_answers(self):
        response = self.post(csv_file(GOOD))
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["total"], 3)
        prime = Question.objects.get(text="Pick the prime number.")
        self.assertEqual(prime.marks, 2)
        self.assertEqual([c.text for c in prime.choices.all()], ["4", "6", "7", "9"])
        self.assertEqual(prime.choices.get(is_correct=True).text, "7")
        self.assertEqual(Question.objects.get(text="The earth orbits the sun.").choices.count(), 2)
        self.assertTrue(AuditLog.objects.filter(action="exams.questions_upload").exists())

    def test_excel_and_alternative_headers(self):
        header = ["Question", "A", "B", "C", "D", "Correct answer", "Mark"]
        response = self.post(xlsx_file([["Largest planet?", "Mars", "Jupiter", "Venus", "Earth", "b", 3]], header))
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Question.objects.get().choices.get(is_correct=True).text, "Jupiter")

    def test_every_problem_is_reported_and_nothing_saved(self):
        rows = [
            GOOD[0],
            ["", "a", "b", "", "", "A", ""],  # no question
            ["Only one option?", "yes", "", "", "", "A", ""],
            ["Gap in options?", "a", "", "c", "", "A", ""],
            ["Answer out of range?", "a", "b", "", "", "D", ""],
            ["Answer text not an option?", "a", "b", "", "", "maybe", ""],
            ["Same options?", "yes", "Yes", "", "", "A", ""],
            ["Bad marks?", "a", "b", "", "", "A", "0"],
            ["What is 2 + 2?", "4", "5", "", "", "A", ""],  # repeats row 2
        ]
        response = self.post(csv_file(rows))
        self.assertEqual(response.status_code, 200)
        problems = {e["row"]: e["errors"] for e in response.data["errors"]}
        self.assertEqual(sorted(problems), [3, 4, 5, 6, 7, 8, 9, 10])
        self.assertIn("question", problems[3])
        self.assertIn("at least two options", problems[4]["options"][0])
        self.assertIn("Option B is empty", problems[5]["options"][0])
        self.assertIn("no option D", problems[6]["answer"][0])
        self.assertIn("exact text", problems[7]["answer"][0])
        self.assertIn("the same", problems[8]["options"][0])
        self.assertIn("marks", problems[9])
        self.assertIn("Same question as row 2", problems[10]["question"][0])
        self.assertFalse(Question.objects.exists())

    def test_missing_columns_and_wrong_file_type(self):
        response = self.post(csv_file([["Q?", "a"]], header=["question", "option_a"]))
        self.assertIn("Missing columns: option_b, answer", str(response.data))
        response = self.post(SimpleUploadedFile("questions.txt", b"hello"))
        self.assertIn(".csv or .xlsx", str(response.data))

    def test_add_to_or_replace_existing_questions(self):
        self.post(csv_file(GOOD[:1]))
        again = self.post(csv_file(GOOD[:1]), dry_run="true")
        self.assertIn("already in the exam", str(again.data["errors"]))
        response = self.post(csv_file(GOOD[1:]))
        self.assertEqual(response.data["total"], 3)
        self.assertEqual(list(Question.objects.values_list("order", flat=True)), [1, 2, 3])
        response = self.post(csv_file(GOOD[:1]), replace="true")
        self.assertEqual((response.data["total"], response.data["replacing"]), (1, 3))

    def test_only_the_course_lecturer_before_the_exam_starts(self):
        self.assertEqual(client_for(self.officer).post(self.url, {"file": csv_file(GOOD)}).status_code, 403)
        self.assertEqual(
            client_for(make_staff(roles=["lecturer"])).post(self.url, {"file": csv_file(GOOD)}).status_code, 404
        )
        start = timezone.localtime() - timedelta(minutes=1)
        self.exam.date, self.exam.start_time = start.date(), start.time()
        self.exam.save()
        response = self.post(csv_file(GOOD))
        self.assertIn("can no longer be changed", str(response.data))

    def test_template_round_trip(self):
        url = f"/api/exams/timetable/{self.exam.pk}/questions/template/"
        blank = self.client_.get(url)
        self.assertEqual(blank.status_code, 200)
        self.assertIn(b"first-in, first-out", blank.content)  # examples when there are no questions yet
        self.post(csv_file(GOOD))
        sheet = load_workbook(BytesIO(self.client_.get(url, {"file": "xlsx"}).content)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[1][0], "What is 2 + 2?")
        self.assertEqual(rows[3][-2:], ("C", 2))
        # The export uploads cleanly as a replacement.
        exported = SimpleUploadedFile("export.xlsx", self.client_.get(url, {"file": "xlsx"}).content)
        response = self.post(exported, replace="true")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["total"], 3)
