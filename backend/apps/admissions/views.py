import csv
from io import BytesIO

from django.db.models import Count, Exists, FloatField, OuterRef, Q
from django.db.models.functions import Cast
from django.http import FileResponse, HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework import mixins, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import NotFound, ValidationError
from rest_framework.parsers import MultiPartParser
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.throttling import ScopedRateThrottle
from rest_framework.views import APIView
from rest_framework_simplejwt.tokens import RefreshToken

from apps.accounts.permissions import IsApplicant, Requires
from apps.core.services import audit
from apps.finance.views.gateway import payment_gateway

from . import payments, services
from .documents import admission_letter_pdf, fee_receipt_pdf
from .models import AdmissionCycle, Application, ApplicationDocument, ApplicationPayment
from .serializers import (
    ApplicantSignupSerializer,
    ApplicationListSerializer,
    ApplicationSerializer,
    CycleSerializer,
    DecisionSerializer,
    DocumentReviewSerializer,
    DocumentSerializer,
    DocumentUploadSerializer,
    EventSerializer,
    OfficerApplicationSerializer,
    PaymentSerializer,
    ScreeningSerializer,
    applicant_payload,
)

S = Application.Status
MANAGE = Requires("admissions.manage")


def pdf_response(content, filename, inline=False):
    response = HttpResponse(content, content_type="application/pdf")
    disposition = "inline" if inline else "attachment"
    response["Content-Disposition"] = f'{disposition}; filename="{filename}"'
    return response


def safe_name(value):
    return value.replace("/", "-")


# --- Public -------------------------------------------------------------------------------------


class CurrentCycleView(APIView):
    """The admission exercise that is currently running (or null), for the website and sign-up."""

    authentication_classes = []
    permission_classes = [AllowAny]

    def get(self, request):
        cycle = AdmissionCycle.current()
        return Response({"cycle": CycleSerializer(cycle).data if cycle else None})


class SignupView(APIView):
    """Create an applicant account and sign straight in."""

    authentication_classes = []
    permission_classes = [AllowAny]
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = "apply"

    def post(self, request):
        serializer = ApplicantSignupSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        user = serializer.save()
        audit(request, "admissions.account_created", user, f"Applicant account created: {user.email}", actor=user)
        refresh = RefreshToken.for_user(user)
        return Response({"access": str(refresh.access_token), "refresh": str(refresh)}, status=status.HTTP_201_CREATED)


# --- Applicant ----------------------------------------------------------------------------------


def my_application(user, cycle=None):
    """The applicant's application for the current cycle, or else their most recent one."""
    qs = Application.objects.filter(applicant=user).select_related(
        "cycle", "programme", "second_choice", "admitted_programme", "applicant"
    )
    cycle = cycle or AdmissionCycle.current()
    return (qs.filter(cycle=cycle).first() if cycle else None) or qs.first()


def require_application(user):
    application = my_application(user)
    if not application:
        raise NotFound("You haven't started an application yet.")
    return application


class MyApplicationView(APIView):
    permission_classes = [IsApplicant]

    def get(self, request):
        cycle = AdmissionCycle.current()
        application = my_application(request.user, cycle)
        return Response(
            {
                "cycle": CycleSerializer(cycle).data if cycle else None,
                "gateway": payment_gateway(),
                **(applicant_payload(application, request) if application else {"application": None}),
            }
        )

    def post(self, request):
        application = services.start_application(request.user, AdmissionCycle.current())
        return Response(applicant_payload(application, request), status=status.HTTP_201_CREATED)

    def patch(self, request):
        application = require_application(request.user)
        serializer = ApplicationSerializer(application, data=request.data, partial=True, context={"request": request})
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(applicant_payload(application, request))


class MyDocumentsView(APIView):
    permission_classes = [IsApplicant]
    parser_classes = [MultiPartParser]

    def post(self, request):
        application = require_application(request.user)
        serializer = DocumentUploadSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        document = services.save_document(
            application,
            serializer.validated_data["kind"],
            serializer.validated_data["file"],
            serializer.content_type,
            request.user,
        )
        return Response(DocumentSerializer(document).data, status=status.HTTP_201_CREATED)


class MyDocumentDetailView(APIView):
    permission_classes = [IsApplicant]

    def delete(self, request, pk):
        document = get_object_or_404(ApplicationDocument, pk=pk, application__applicant=request.user)
        services.delete_document(document)
        return Response(status=status.HTTP_204_NO_CONTENT)


class MyPaymentView(APIView):
    """Start paying the application fee: returns the Paystack checkout URL."""

    permission_classes = [IsApplicant]

    def post(self, request):
        application = require_application(request.user)
        if not request.user.email:
            raise ValidationError("Add an email address to your account before paying.")
        callback = request.build_absolute_uri("/portal/application?step=payment")
        payment, url = payments.start(application, callback)
        return Response({"authorization_url": url, "payment": PaymentSerializer(payment).data})


class MyPaymentVerifyView(APIView):
    permission_classes = [IsApplicant]

    def post(self, request):
        reference = str(request.data.get("reference", ""))
        payment = get_object_or_404(ApplicationPayment, reference=reference, application__applicant=request.user)
        payment = payments.verify(payment)
        return Response(PaymentSerializer(payment).data)


class MySubmitView(APIView):
    permission_classes = [IsApplicant]

    def post(self, request):
        application = services.submit(require_application(request.user), request.user, request)
        return Response(applicant_payload(application, request))


class MyAcceptView(APIView):
    permission_classes = [IsApplicant]

    def post(self, request):
        application = services.accept_offer(require_application(request.user), request.user, request)
        return Response(applicant_payload(application, request))


class MyLetterView(APIView):
    permission_classes = [IsApplicant]

    def get(self, request):
        application = require_application(request.user)
        if application.status not in (S.ADMITTED, S.ACCEPTED):
            raise NotFound("Your admission letter isn't available.")
        return pdf_response(
            admission_letter_pdf(application), f"admission-letter-{safe_name(application.letter_number)}.pdf"
        )


class MyReceiptView(APIView):
    permission_classes = [IsApplicant]

    def get(self, request):
        payment = (
            ApplicationPayment.objects.filter(
                application__applicant=request.user, status=ApplicationPayment.Status.SUCCESS
            )
            .select_related("application__cycle", "application__applicant")
            .first()
        )
        if not payment:
            raise NotFound("You haven't paid the application fee yet.")
        return pdf_response(fee_receipt_pdf(payment), f"application-fee-receipt-{payment.reference}.pdf")


class DocumentFileView(APIView):
    """An uploaded document, for its applicant and the Admissions Office only."""

    def get(self, request, pk):
        document = get_object_or_404(ApplicationDocument.objects.select_related("application"), pk=pk)
        if document.application.applicant_id != request.user.id and not request.user.has_permission(
            "admissions.manage"
        ):
            raise NotFound()
        return FileResponse(
            document.file.open("rb"), content_type=document.content_type, filename=document.original_filename
        )


# --- Admissions Office --------------------------------------------------------------------------

AGGREGATE = Cast("utme_score", FloatField()) / 8 + Cast("screening_score", FloatField()) / 2
EXPORT_COLUMNS = [
    ("Application no.", lambda a: a.number),
    ("Surname", lambda a: a.applicant.last_name),
    ("First name", lambda a: a.applicant.first_name),
    ("Middle name", lambda a: a.middle_name),
    ("Email", lambda a: a.applicant.email),
    ("Phone", lambda a: a.phone),
    ("Gender", lambda a: a.get_gender_display()),
    ("State of origin", lambda a: a.state_of_origin),
    ("Entry mode", lambda a: a.get_entry_mode_display()),
    ("First choice", lambda a: a.programme.title if a.programme else ""),
    ("Second choice", lambda a: a.second_choice.title if a.second_choice else ""),
    ("Faculty", lambda a: a.programme.department.faculty.name if a.programme else ""),
    ("JAMB reg. no.", lambda a: a.jamb_reg_number),
    ("UTME score", lambda a: a.utme_score),
    ("Screening score", lambda a: float(a.screening_score) if a.screening_score is not None else None),
    ("Aggregate", lambda a: a.aggregate_score),
    ("Status", lambda a: a.get_status_display()),
    ("Fee paid", lambda a: "Yes" if a.paid else "No"),
    ("Submitted", lambda a: timezone.localtime(a.submitted_at).strftime("%Y-%m-%d %H:%M") if a.submitted_at else ""),
    ("Admitted into", lambda a: a.admitted_programme.title if a.admitted_programme else ""),
    ("Letter no.", lambda a: a.letter_number or ""),
]


def allowed_actions(application):
    """What the Admissions Office can do next, for the review screen's buttons."""
    required_verified = all(
        any(d.kind == k and d.status == ApplicationDocument.Status.VERIFIED for d in application.documents.all())
        for k in ApplicationDocument.REQUIRED
    )
    scored = application.screening_score is not None
    return {
        S.SUBMITTED: ["start_review", "review_documents"],
        S.UNDER_REVIEW: ["review_documents", "reject"] + (["to_screening"] if required_verified else []),
        S.SCREENING: ["record_screening", "reject"] + (["approve", "waitlist"] if scored else []),
        S.WAITLISTED: ["approve", "reject"],
        S.APPROVED: ["issue_letter"],
        S.ADMITTED: ["download_letter"],
        S.ACCEPTED: ["download_letter"],
    }.get(application.status, [])


class ApplicationViewSet(viewsets.ReadOnlyModelViewSet):
    """Applications for the Admissions Office: search, filter, review and decide."""

    permission_classes = [MANAGE]
    serializer_class = ApplicationListSerializer
    filterset_fields = {
        "status": ["exact"],
        "cycle": ["exact"],
        "entry_mode": ["exact"],
        "programme": ["exact"],
        "programme__department": ["exact"],
        "programme__department__faculty": ["exact"],
        "state_of_origin": ["exact"],
    }
    search_fields = [
        "number",
        "applicant__first_name",
        "applicant__last_name",
        "applicant__email",
        "jamb_reg_number",
        "phone",
    ]
    ordering_fields = ["submitted_at", "utme_score", "screening_score", "aggregate", "number"]
    ordering = ["-submitted_at", "-created_at"]

    def get_queryset(self):
        qs = (
            Application.objects.select_related(
                "applicant", "cycle", "programme__department__faculty", "second_choice", "admitted_programme"
            )
            .prefetch_related("documents")
            .annotate(
                paid=Exists(ApplicationPayment.objects.filter(application=OuterRef("pk"), status="success")),
                documents_pending=Count("documents", filter=Q(documents__status=ApplicationDocument.Status.PENDING)),
                aggregate=AGGREGATE,
            )
        )
        params = self.request.query_params
        if self.action in ("list", "export") and not params.get("status"):
            qs = qs.exclude(status=S.DRAFT)  # unfinished drafts only when asked for
        if params.get("fee_paid") in ("true", "false"):
            qs = qs.filter(paid=params["fee_paid"] == "true")
        if params.get("cycle") is None and self.action in ("list", "export", "summary"):
            cycle = AdmissionCycle.current()
            if cycle:
                qs = qs.filter(cycle=cycle)
        return qs

    def retrieve(self, request, *args, **kwargs):
        application = self.get_object()
        return Response(self.detail_payload(application))

    def detail_payload(self, application):
        application = Application.objects.select_related(
            "applicant",
            "cycle",
            "programme",
            "second_choice",
            "admitted_programme",
            "reviewer",
            "screened_by",
            "decided_by",
        ).get(pk=application.pk)
        return {
            "application": OfficerApplicationSerializer(application).data,
            "documents": DocumentSerializer(application.documents.select_related("reviewed_by"), many=True).data,
            "required_documents": list(ApplicationDocument.REQUIRED),
            "payments": PaymentSerializer(application.payments.all(), many=True).data,
            "timeline": EventSerializer(application.events.select_related("actor", "application"), many=True).data,
            "checklist": services.checklist(application),
            "actions": allowed_actions(application),
            "olevel_problems": services.olevel_problems(application.olevel_results),
        }

    def _act(self, request, fn, *args, **kwargs):
        application = fn(self.get_object(), *args, by=request.user, request=request, **kwargs)
        return Response(self.detail_payload(application))

    @action(detail=True, methods=["post"], url_path="start-review")
    def start_review(self, request, pk=None):
        return self._act(request, services.start_review)

    @action(detail=True, methods=["post"], url_path="to-screening")
    def to_screening(self, request, pk=None):
        return self._act(request, services.move_to_screening)

    @action(detail=True, methods=["post"])
    def screening(self, request, pk=None):
        serializer = ScreeningSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        return self._act(request, services.record_screening, data["score"], data["remarks"].strip())

    @action(detail=True, methods=["post"])
    def decide(self, request, pk=None):
        serializer = DecisionSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        return self._act(
            request, services.decide, data["decision"], data["note"].strip(), programme=data.get("programme")
        )

    @action(detail=True, methods=["post"], url_path="issue-letter")
    def issue_letter(self, request, pk=None):
        return self._act(request, services.issue_letter)

    @action(detail=True, methods=["post"])
    def note(self, request, pk=None):
        """An internal note on the timeline (not shown to the applicant)."""
        text = str(request.data.get("note", "")).strip()
        if not text:
            raise ValidationError({"note": "Write a note."})
        application = self.get_object()
        services.record(application, "note", request.user, text)
        audit(request, "admissions.note", application, f"{application.number}: note added")
        return Response(self.detail_payload(application))

    @action(detail=True, methods=["get"])
    def letter(self, request, pk=None):
        application = self.get_object()
        if not application.letter_number:
            raise NotFound("No admission letter has been issued for this application.")
        inline = request.query_params.get("inline") == "1"
        return pdf_response(
            admission_letter_pdf(application), f"admission-letter-{safe_name(application.letter_number)}.pdf", inline
        )

    @action(detail=False, methods=["get"])
    def summary(self, request):
        qs = self.filter_queryset(self.get_queryset()).order_by()
        # Count from a plain queryset: the list's document-count join would multiply the rows.
        plain = Application.objects.filter(pk__in=qs.values("pk")).order_by()
        counts = dict(plain.values_list("status").annotate(n=Count("id")))
        cycle = AdmissionCycle.current()
        return Response(
            {
                "cycle": CycleSerializer(cycle).data if cycle else None,
                "by_status": {s: counts.get(s, 0) for s in S.values},
                "submitted_total": sum(n for s, n in counts.items() if s != S.DRAFT),
                "fees_paid": qs.filter(paid=True).values("pk").distinct().count(),
            }
        )

    @action(detail=False, methods=["get"])
    def export(self, request):
        applications = list(self.filter_queryset(self.get_queryset()))
        stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
        audit(request, "admissions.export", None, f"Exported {len(applications)} applications")
        header = [label for label, _ in EXPORT_COLUMNS]
        rows = [[get(a) for _, get in EXPORT_COLUMNS] for a in applications]
        if request.query_params.get("file") == "xlsx":
            book = Workbook()
            sheet = book.active
            sheet.title = "Applications"
            sheet.append(header)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="12305A")
            for row in rows:
                sheet.append(row)
            for column in sheet.columns:
                width = max(len(str(c.value or "")) for c in column)
                sheet.column_dimensions[column[0].column_letter].width = min(max(10, width + 2), 40)
            sheet.freeze_panes = "B2"
            buffer = BytesIO()
            book.save(buffer)
            response = HttpResponse(
                buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            response["Content-Disposition"] = f'attachment; filename="applications-{stamp}.xlsx"'
            return response
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="applications-{stamp}.csv"'
        response.write("﻿")  # BOM so Excel opens UTF-8 names correctly
        writer = csv.writer(response)
        writer.writerow(header)
        writer.writerows(["" if v is None else v for v in row] for row in rows)
        return response


class DocumentReviewView(APIView):
    permission_classes = [MANAGE]

    def post(self, request, pk):
        document = get_object_or_404(ApplicationDocument.objects.select_related("application"), pk=pk)
        serializer = DocumentReviewSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        services.review_document(
            document,
            serializer.validated_data["verified"],
            serializer.validated_data["note"].strip(),
            request.user,
            request,
        )
        return Response(ApplicationViewSet().detail_payload(document.application))


class CycleViewSet(mixins.ListModelMixin, mixins.CreateModelMixin, mixins.UpdateModelMixin, viewsets.GenericViewSet):
    permission_classes = [MANAGE]
    serializer_class = CycleSerializer
    queryset = AdmissionCycle.objects.all()
    pagination_class = None

    def perform_create(self, serializer):
        cycle = serializer.save()
        audit(self.request, "admissions.cycle_create", cycle, f"Created {cycle}")

    def perform_update(self, serializer):
        cycle = serializer.save()
        audit(self.request, "admissions.cycle_update", cycle, f"Updated {cycle}")
